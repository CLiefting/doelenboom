"""Tests voor app/program_pptx.py (DOEL-112) -- de programma-PowerPoint voor
meerdere projecten. Controleert de slide-opbouw, de indeling van de
deliverables, de navigatie (interne links + Home-knop) en de invoer-
validatie door het resultaat weer in te laden met python-pptx."""
from __future__ import annotations

import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

from app.main import app
from app.program_pptx import (
    MAX_PROJECTS,
    ProgramPptxError,
    add_months_same_day,
    build_program_pptx,
    classify_deliverables,
    paginate_deliverables,
    parse_options,
    section_titles,
    sort_projects,
)

REF = date(2026, 10, 13)
META = {'doelenboom': 'OVpay @ NS', 'tenant': 'NS', 'exportedAt': '2026-10-13T18:00:00Z', 'exportedBy': 'tester'}


def prod(i, name, verwacht=None, werkelijk=None, typ='deliverable', deadline=None, pct=0):
    return {'id': i, 'name': name, 'type': typ, 'verwachteDatum': verwacht, 'werkelijkeDatum': werkelijk,
            'deadline': deadline, 'pctGereed': pct}


def act(i, name, start, end, summary=False, milestone=False, wbs=None):
    return {'id': i, 'name': name, 'startDate': start, 'endDate': end, 'isSummary': summary,
            'isMilestone': milestone, 'wbs': wbs}


def proj(code, name, products=(), activities=(), goals=None, rag='groen'):
    return {
        'project': {'code': code, 'name': name, 'description': '', 'status': {'rag': rag}, 'tags': [], 'orgs': []},
        'products': list(products), 'activities': list(activities),
        'productDependencies': [], 'activityDependencies': [], 'goals': goals or {},
    }


def sample_projects():
    return [
        proj('NP37', 'Sweepen',
             [prod(1, 'Ontwerp', '2026-09-20', '2026-09-25', pct=100), prod(2, 'Go-live', '2026-12-15', typ='mijlpaal'),
              prod(3, 'Testrapport', '2026-11-10'), prod(4, 'Handleiding')],
             [act(1, 'Ontwerp', '2026-07-01', '2026-09-30', True, wbs='1'),
              act(2, 'Realisatie', '2026-10-01', '2026-12-15', True, wbs='2'),
              act(3, 'Deeltaak', '2026-10-01', '2026-10-15', True, wbs='2.1')],
             {'direct': ['Automatisch sweepen'], 'top': ['Eenvoudig betalen']}),
        proj('NP12', 'Check-in check-out', [prod(5, 'Pilot', '2026-09-30'), prod(6, 'Besluit', '2027-03-01', typ='mijlpaal')],
             rag='oranje'),
        proj('NP05', 'Aanpassing back-office', [prod(8, 'Koppeling', '2026-10-01', deadline='2026-10-10')], rag='rood'),
    ]


def build(projects=None, options=None, goal_columns=None):
    data = {'projects': projects if projects is not None else sample_projects(),
            'goalColumns': goal_columns if goal_columns is not None else {'direct': 'Capability', 'top': 'Strategisch doel'}}
    content = build_program_pptx(data, options if options is not None else {'referenceDate': '2026-10-13', 'windowMonths': 3}, META)
    return Presentation(io.BytesIO(content))


def slide_text(slide) -> str:
    return '\n'.join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def jump_targets(prs, slide) -> list[int]:
    """Paginanummers (1-based) waar klikbare vormen op deze slide heen springen."""
    index = {s.part.partname: i + 1 for i, s in enumerate(prs.slides)}
    targets = []
    for shape in slide.shapes:
        h = shape._element.find('.//' + qn('a:hlinkClick'))
        if h is None or h.get('action') != 'ppaction://hlinksldjump':
            continue
        rel = slide.part.rels[h.get(qn('r:id'))]
        targets.append(index[rel.target_part.partname])
    return targets


class TestHulpfuncties:
    def test_maanden_optellen_klemt_op_maandeinde(self):
        assert add_months_same_day(date(2026, 10, 13), 3) == date(2027, 1, 13)
        assert add_months_same_day(date(2027, 1, 31), 1) == date(2027, 2, 28)
        assert add_months_same_day(date(2028, 1, 31), 1) == date(2028, 2, 29)

    def test_projecten_alfabetisch(self):
        names = [p['project']['name'] for p in sort_projects(sample_projects())]
        assert names == ['Aanpassing back-office', 'Check-in check-out', 'Sweepen']

    def test_standaard_peildatum_is_exportdatum_en_venster_3(self):
        assert parse_options({}, META) == (REF, 3)

    @pytest.mark.parametrize('options', [
        {'referenceDate': '13-10-2026'}, {'referenceDate': '2026-02-30'}, {'referenceDate': 20261013},
        {'referenceDate': '1899-01-01'}, {'windowMonths': 0}, {'windowMonths': 13}, {'windowMonths': '3'},
        {'windowMonths': True},
    ])
    def test_ongeldige_keuzes(self, options):
        with pytest.raises(ProgramPptxError):
            parse_options(options, META)


class TestDeliverables:
    def test_indeling_in_blokken(self):
        s = classify_deliverables(sort_projects(sample_projects()), REF, 3)
        assert [r['name'] for r in s['late']] == ['Pilot', 'Koppeling']
        assert [r['name'] for r in s['window']] == ['Testrapport', 'Go-live']
        assert [r['name'] for r in s['later']] == ['Besluit']
        assert [r['name'] for r in s['undated']] == ['Handleiding']
        # Opgeleverd telt nergens mee.
        assert all(r['name'] != 'Ontwerp' for rows in s.values() for r in rows)

    def test_venstergrens_is_exclusief(self):
        p = [proj('A', 'A', [prod(1, 'op peildatum', '2026-10-13'), prod(2, 'laatste dag', '2027-01-12'),
                              prod(3, 'eerste dag later', '2027-01-13')])]
        s = classify_deliverables(p, REF, 3)
        assert [r['name'] for r in s['window']] == ['op peildatum', 'laatste dag']
        assert [r['name'] for r in s['later']] == ['eerste dag later']

    def test_venster_instelbaar(self):
        s = classify_deliverables(sort_projects(sample_projects()), REF, 1)
        assert [r['name'] for r in s['window']] == ['Testrapport']

    def test_titels(self):
        t = section_titles(REF, 3)
        assert t['window'] == 'Komende 3 maanden — 13 okt 2026 – 12 jan 2027'
        assert 'vóór 13 okt 2026' in t['late']
        assert section_titles(REF, 1)['window'].startswith('Komende 1 maand ')

    def test_paginering_herhaalt_bloktitel_met_vervolg(self):
        rows = [{'name': str(i)} for i in range(20)]
        pages = paginate_deliverables({'late': [], 'window': rows, 'later': [], 'undated': []}, rows_per_slide=14)
        assert len(pages) == 2
        assert pages[0][0] == ('section', 'window', False)
        assert pages[1][0] == ('section', 'window', True)
        assert sum(1 for p in pages for e in p if e[0] == 'row') == 20

    def test_bloktitel_nooit_als_laatste_regel(self):
        pages = paginate_deliverables(
            {'late': [{'name': str(i)} for i in range(13)], 'window': [{'name': 'x'}], 'later': [], 'undated': []},
            rows_per_slide=14,
        )
        assert pages[0][-1][0] == 'row'
        assert pages[1][0] == ('section', 'window', False)


class TestPresentatie:
    def test_opbouw_en_volgorde(self):
        prs = build()
        texts = [slide_text(s) for s in prs.slides]
        assert len(texts) == 6  # tijdlijn, deliverables, doelen, 3x voortgang
        assert 'Tijdlijnen' in texts[0]
        assert 'Deliverables' in texts[1]
        assert 'Bijdrage aan doelen' in texts[2]
        # Voortgangsslides alfabetisch: de projectnaam is de grote titel.
        def title(slide):
            return next(sh.text_frame.text for sh in slide.shapes
                        if sh.has_text_frame and sh.text_frame.paragraphs[0].runs
                        and sh.text_frame.paragraphs[0].runs[0].font.size == Pt(28))
        slides = list(prs.slides)
        assert [title(s) for s in slides[3:]] == ['Aanpassing back-office', 'Check-in check-out', 'Sweepen']

    def test_tijdlijn_toont_fasen_van_bovenste_niveau_en_peildatum(self):
        t = slide_text(build().slides[0])
        assert 'Ontwerp' in t and 'Realisatie' in t
        assert 'Deeltaak' not in t
        assert 'vandaag' in t
        t2 = slide_text(build(options={'referenceDate': '2026-11-01'}).slides[0])
        assert 'peildatum' in t2

    def test_deliverables_slide_te_laat_bovenaan(self):
        t = slide_text(build().slides[1])
        assert t.index('Te laat') < t.index('Komende 3 maanden') < t.index('Later') < t.index('Nog geen datum')
        assert '2 te laat' in t

    def test_doelen(self):
        t = slide_text(build().slides[2])
        assert 'Capability' in t and 'Strategisch doel' in t
        assert 'Automatisch sweepen' in t and 'Eenvoudig betalen' in t

    def test_doelen_zonder_bovenste_kolom(self):
        t = slide_text(build(goal_columns={'direct': 'Capability'}).slides[2])
        assert 'Strategisch doel' not in t

    def test_links_naar_voortgang_en_home(self):
        prs = build()
        slides = list(prs.slides)
        # Tijdlijn: projectnamen/balken springen naar slide 4-6, geen Home.
        assert set(jump_targets(prs, slides[0])) == {4, 5, 6}
        # Elke andere slide heeft een Home-knop naar slide 1.
        for s in slides[1:]:
            assert any(sh.name == 'home-knop' for sh in s.shapes)
            assert 1 in jump_targets(prs, s)
        assert not any(sh.name == 'home-knop' for sh in slides[0].shapes)
        # Deliverables springen naar de voortgang van hun project.
        assert {4, 5, 6} <= set(jump_targets(prs, slides[1]))

    def test_veel_projecten_pagineert_tijdlijn(self):
        projects = [proj(f'P{i:02d}', f'Project {i:02d}', [prod(i, 'd', '2026-11-01')]) for i in range(20)]
        prs = build(projects=projects)
        texts = [slide_text(s) for s in prs.slides]
        assert sum(1 for t in texts if 'Tijdlijnen' in t) == 3
        assert 'pagina 1/3' in texts[0]

    def test_geen_openstaande_deliverables(self):
        t = slide_text(build(projects=[proj('A', 'A', [prod(1, 'klaar', '2026-01-01', '2026-01-01')])]).slides[1])
        assert 'geen openstaande deliverables' in t

    def test_project_zonder_datums(self):
        assert 'Nog geen datums vastgelegd' in slide_text(build(projects=[proj('A', 'A')]).slides[0])


class TestVeiligheid:
    def test_besturingstekens_en_html_komen_als_platte_tekst(self):
        evil = '<script>alert(1)</script>\x00\x07"\'&'
        prs = build(projects=[proj('X', evil, [prod(1, evil, '2026-11-01')], goals={'direct': [evil]})])
        t = '\n'.join(slide_text(s) for s in prs.slides)
        assert '<script>alert(1)</script>' in t
        assert '\x00' not in t and '\x07' not in t

    def test_te_veel_projecten(self):
        with pytest.raises(ProgramPptxError):
            build(projects=[proj(f'P{i}', f'P{i}') for i in range(MAX_PROJECTS + 1)])

    @pytest.mark.parametrize('data', [None, {}, {'projects': []}, {'projects': 'x'}, {'projects': [1]},
                                      {'projects': [{'project': 'x'}]}])
    def test_ongeldige_gegevens(self, data):
        with pytest.raises(ProgramPptxError):
            build_program_pptx(data, {}, META)

    def test_te_veel_slides(self):
        many = [prod(i, f'd{i}', '2026-11-01') for i in range(5000)]
        with pytest.raises(ProgramPptxError, match='maximum'):
            build(projects=[proj('A', 'A', many)])


class TestEndpoint:
    client = TestClient(app)

    def test_levert_pptx(self):
        res = self.client.post('/program-pptx', json={
            'data': {'projects': sample_projects(), 'goalColumns': {'direct': 'Capability'}},
            'options': {'windowMonths': 3}, 'meta': META,
        })
        assert res.status_code == 200
        assert res.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.presentationml')
        assert len(Presentation(io.BytesIO(res.content)).slides) == 6

    def test_ongeldige_keuze_geeft_422(self):
        res = self.client.post('/program-pptx', json={
            'data': {'projects': sample_projects()}, 'options': {'windowMonths': 99}, 'meta': META,
        })
        assert res.status_code == 422
        assert 'venster' in res.json()['error']
