"""Tests voor app/tree_pptx.py (DOEL-88) -- bouwt de PowerPoint van een hele
doelenboom en controleert structuur en inhoud door het resultaat weer in te
laden met python-pptx, plus het /tree-pptx-eindpunt in main.py."""
from __future__ import annotations

import io
import re
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.action import PP_ACTION
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from pptx.util import Inches

from app.main import app
from app.project_pptx import CONTENT_W, MARGIN, SLIDE_H
from app.tree_pptx import (
    CHAIN_MAX_ROWS,
    MAX_SLIDES,
    MAX_SLIDES_PER_ELEMENT,
    OVERVIEW_MAX_LINES,
    OVERVIEW_MAX_ROWS,
    OVERVIEW_NODES_BOTTOM,
    OVERVIEW_NODES_TOP,
    RIGHT_LEFT,
    SECTION_MAX_CHIPS,
    TreePptxError,
    _clamp_per_row,
    _Model,
    _paginate_blocks,
    _rgb,
    _text_on,
    _tint,
    build_tree_pptx,
)

client = TestClient(app)
META = {'doelenboom': 'Bewaken en Beveiligen', 'tenant': 'CLPK', 'exportedAt': '2026-10-05T10:00:00Z'}


def col(position, type_name, color='#3E6FA6', label=None, **extra):
    return {
        'position': position, 'typeName': type_name, 'title': extra.pop('title', type_name),
        'subtitle': extra.pop('subtitle', f'Vraag bij {type_name}?'), 'color': color,
        'relationLabelToNext': label, 'aliases': extra.pop('aliases', []),
    }


def el(code, type_name, name=None, **extra):
    return {'code': code, 'type': type_name, 'name': name or f'Naam {code}', 'description': '', 'kpi': '',
            'taakveld': '', 'subtaakveld': '', **extra}


def edge(source, target, weight='primair'):
    return {'source': source, 'target': target, 'weight': weight, 'toelichting': 'GEHEIM-TOELICHTING'}


def make_data():
    """Project -> Capability -> Benefit -> Doel -> Missie."""
    return {
        'columns': [
            col(0, 'Project', '#3E6FA6', 'ontwikkelt'),
            col(1, 'Capability', '#6B4C8A', 'ondersteunt'),
            col(2, 'Benefit', '#C05A2C', 'realiseert'),
            col(3, 'Doel', '#8FAADC', 'geeft invulling aan'),
            col(4, 'Missie', '#203864', None),
        ],
        'elements': [
            el('M1', 'Missie', 'Veilig land'),
            el('D1', 'Doel', 'Betrouwbare beveiliging'),
            el('D2', 'Doel', 'Informatiegestuurd werken'),
            el('B2', 'Benefit', 'Tweede benefit'),
            el('B1', 'Benefit', 'Eerste benefit', description='Regel een.\nRegel twee.', kpi='Reactietijd < 5 min',
               taakveld='Bewaken', subtaakveld='Objecten'),
            el('C1', 'Capability', 'Meldkamer'),
            el('C2', 'Capability', 'Sensoren'),
            el('P1', 'Project', 'Nieuw systeem'),
            el('P2', 'Project', 'Los project'),
        ],
        'edges': [
            edge('D1', 'M1'), edge('D2', 'M1'),
            edge('B1', 'D1'), edge('B1', 'D2', 'ondersteunend'), edge('B2', 'D2'),
            edge('C1', 'B1'), edge('C2', 'B2'),
            edge('P1', 'C1'),
        ],
        'tags': [{'code': 'T1', 'name': 'Digitalisering'}],
        'elementTags': {'B1': ['T1']},
        'orgUnits': [{'code': 'O1', 'name': 'Brigade Noord'}],
        'obOrg': {'B1': [{'org': 'O1', 'relatietype': 'eigenaar'}]},
    }


def load(content: bytes) -> Presentation:
    return Presentation(io.BytesIO(content))


def all_text(slide) -> str:
    return '\n'.join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def full_text(prs) -> str:
    return '\n'.join(all_text(s) for s in prs.slides)


def shape_names(slide) -> list[str]:
    return [s.name for s in slide.shapes]


def build(options=None, data=None):
    return load(build_tree_pptx(data or make_data(), options or {}, META))


class TestHelpers:
    def test_per_row_wordt_begrensd(self):
        assert _clamp_per_row(4) == 4
        assert _clamp_per_row(1) == 2
        assert _clamp_per_row(99) == 6
        assert _clamp_per_row('3') == 3
        assert _clamp_per_row(None) == 4
        assert _clamp_per_row('veel') == 4
        assert _clamp_per_row({'a': 1}) == 4

    def test_ongeldige_kleur_valt_terug_op_grijs(self):
        assert str(_rgb('#3E6FA6')) == '3E6FA6'
        for bad in (None, '', 'rood', '#12', '#GGGGGG', 123, '"/><script>'):
            assert str(_rgb(bad)) == '6C6F76'

    def test_tekstkleur_past_bij_achtergrond(self):
        assert str(_text_on(_rgb('#203864'))) == 'FFFFFF'
        assert str(_text_on(_rgb('#F4F5F7'))) != 'FFFFFF'

    def test_paginering_verdeelt_lange_tekst_en_kapt_op_het_eind_af(self):
        blocks = [('Beschrijving', ['woord ' * 400])]
        pages = _paginate_blocks(blocks, 40, [5, 10, 10])
        assert len(pages) == 3
        assert pages[0][0] == ('h', 'Beschrijving')
        assert pages[1][0] == ('h', 'Beschrijving (vervolg)')
        assert pages[-1][-1][1].endswith('…')

    def test_paginering_korte_tekst_blijft_op_een_pagina(self):
        pages = _paginate_blocks([('A', ['kort']), ('B', ['ook kort'])], 40, [12, 20])
        assert pages == [[('h', 'A'), ('p', 'kort'), ('h', 'B'), ('p', 'ook kort')]]


class TestSnoerSlide:
    def test_alle_kolommen_met_omschrijving_relatielabel_en_aantal(self):
        prs = build({'perRow': 4})
        text = all_text(prs.slides[0])
        for n, name in enumerate(['Project', 'Capability', 'Benefit', 'Doel', 'Missie'], start=1):
            assert f'{n}. {name}' in text
            assert f'Vraag bij {name}?' in text
        for label in ('ontwikkelt', 'ondersteunt', 'realiseert', 'geeft invulling aan'):
            assert label in text
        assert '2 elementen' in text and '1 element' in text
        assert 'Bewaken en Beveiligen' in text

    def test_zonder_gekozen_kolommen_alleen_het_snoer_en_de_totale_boom(self):
        prs = build({})
        assert len(prs.slides) == 2
        assert 'uitgewerkt' not in all_text(prs.slides[0])
        assert 'DE TOTALE BOOM' in all_text(prs.slides[1])

    def test_gekozen_kolom_is_gemarkeerd(self):
        text = all_text(build({'slideColumns': ['Benefit']}).slides[0])
        assert '2 elementen · per element uitgewerkt' in text

    def test_slingert_naar_volgende_regel(self):
        """perRow=2 met 5 kolommen: 3 regels; regel 2 loopt van rechts naar
        links, dus kolom 3 staat rechts van kolom 4."""
        slide = build({'perRow': 2}).slides[0]
        boxes = {s.name: s for s in slide.shapes if s.name.startswith('snoer-kolom-')}
        assert len(boxes) == 5
        assert boxes['snoer-kolom-1'].left < boxes['snoer-kolom-2'].left
        assert boxes['snoer-kolom-3'].top > boxes['snoer-kolom-1'].top
        assert boxes['snoer-kolom-3'].left > boxes['snoer-kolom-4'].left
        assert boxes['snoer-kolom-3'].left == boxes['snoer-kolom-2'].left
        assert boxes['snoer-kolom-5'].top > boxes['snoer-kolom-3'].top
        assert boxes['snoer-kolom-5'].left == boxes['snoer-kolom-4'].left

    def test_vakken_blijven_binnen_de_slide(self):
        for per_row in (2, 3, 4, 5, 6):
            prs = build({'perRow': per_row})
            for shape in prs.slides[0].shapes:
                assert shape.left >= 0 and shape.top >= 0
                assert shape.left + shape.width <= prs.slide_width
                assert shape.top + shape.height <= prs.slide_height

    def test_veel_kolommen_lopen_door_op_een_vervolgslide(self):
        data = {'columns': [col(i, f'K{i}') for i in range(11)], 'elements': [], 'edges': []}
        prs = build({'perRow': 2}, data)  # 6 regels -> 2 slides, daarna de totale boom
        assert len(prs.slides) == 3
        assert 'DE TOTALE BOOM' in all_text(prs.slides[2])
        assert '8. K7' in all_text(prs.slides[0]) and '9. K8' not in all_text(prs.slides[0])
        assert '9. K8' in all_text(prs.slides[1])
        assert 'VERVOLG' in all_text(prs.slides[1])


class TestVerborgenKolommen:
    OPTIONS = {'visibleColumns': ['Project', 'Benefit', 'Doel', 'Missie'], 'slideColumns': ['Benefit', 'Capability']}

    def test_verborgen_kolom_komt_nergens_voor(self):
        prs = build(self.OPTIONS)
        text = full_text(prs)
        assert 'Capability' not in text
        assert 'Meldkamer' not in text and 'Sensoren' not in text
        assert 'C1' not in text

    def test_verborgen_kolom_krijgt_geen_element_slides_ook_al_is_hij_gekozen(self):
        prs = build(self.OPTIONS)
        # snoer + totale boom + tussenslide Benefit + B2 + B1
        assert len(prs.slides) == 5

    def test_relatie_wordt_doorgetrokken_over_de_verborgen_kolom(self):
        model = _Model(make_data(), self.OPTIONS)
        assert model.parents['P1'] == {'B1': True}
        assert model.children['B1'] == {'P1': True}
        prs = build(self.OPTIONS)
        b1 = next(s for s in prs.slides if 'B1 — Eerste benefit' in all_text(s))
        assert 'P1  Nieuw systeem' in all_text(b1)
        assert 'verbinding P1' in shape_names(b1)
        assert 'relatie P1 > B1' in shape_names(b1)

    def test_geen_relatielabel_tussen_kolommen_die_geen_buren_zijn(self):
        text = all_text(build(self.OPTIONS).slides[0])
        assert 'ontwikkelt' not in text  # Project -> (verborgen Capability) -> Benefit
        assert 'realiseert' in text

    def test_doorgetrokken_relatie_is_ondersteunend_als_een_schakel_dat_is(self):
        data = make_data()
        data['edges'] = [edge('P1', 'C1', 'ondersteunend'), edge('C1', 'B1')]
        model = _Model(data, self.OPTIONS)
        assert model.parents['P1'] == {'B1': False}

    def test_lege_selectie_van_zichtbare_kolommen_wordt_geweigerd(self):
        with pytest.raises(TreePptxError):
            build_tree_pptx(make_data(), {'visibleColumns': []}, META)
        with pytest.raises(TreePptxError):
            build_tree_pptx(make_data(), {'visibleColumns': ['Bestaat niet']}, META)


class TestElementSlides:
    OPTIONS = {'slideColumns': ['Benefit']}

    def slide_for(self, prs, title):
        return next(s for s in prs.slides if title in all_text(s))

    def test_tussenslide_en_slide_per_element_in_boomvolgorde(self):
        prs = build(self.OPTIONS)
        assert len(prs.slides) == 5
        assert 'KOLOM' in all_text(prs.slides[2]) and 'Benefit' in all_text(prs.slides[2])
        # Boomvolgorde = volgorde van aanleveren: B2 voor B1.
        assert 'B2 — Tweede benefit' in all_text(prs.slides[3])
        assert 'B1 — Eerste benefit' in all_text(prs.slides[4])

    def test_meerdere_kolommen_in_kolomvolgorde(self):
        prs = build({'slideColumns': ['Missie', 'Project']})
        assert 'KOLOM' in all_text(prs.slides[2]) and 'Project' in all_text(prs.slides[2])
        assert 'P1 — Nieuw systeem' in all_text(prs.slides[3])
        assert 'M1 — Veilig land' in all_text(prs.slides[-1])
        assert len(prs.slides) == 2 + (1 + 2) + (1 + 1)

    def test_hele_pad_omhoog_en_omlaag_maar_geen_zijtakken(self):
        slide = self.slide_for(build(self.OPTIONS), 'B1 — Eerste benefit')
        names = shape_names(slide)
        assert 'element B1 (geselecteerd)' in names
        for code in ('P1', 'C1', 'D1', 'D2', 'M1'):
            assert f'element {code}' in names
        # B2, C2 (zijtak via D2) en het losse P2 horen niet in de keten van B1.
        for code in ('B2', 'C2', 'P2'):
            assert f'element {code}' not in names

    def test_ondersteunende_relatie_is_gestippeld(self):
        slide = self.slide_for(build(self.OPTIONS), 'B1 — Eerste benefit')
        lines = {s.name: s for s in slide.shapes if s.name.startswith('relatie ')}
        assert lines['relatie B1 > D2 (ondersteunend)'].line.dash_style == MSO_LINE.DASH
        assert lines['relatie B1 > D1'].line.dash_style is None
        assert 'stippellijn' in all_text(slide)

    def test_tekst_beschrijving_kpi_taakveld_verbindingen_tags_en_org(self):
        text = all_text(self.slide_for(build(self.OPTIONS), 'B1 — Eerste benefit'))
        assert 'BENEFIT · BEWAKEN / OBJECTEN' in text
        assert 'Regel een.' in text and 'Regel twee.' in text
        assert 'Reactietijd < 5 min' in text
        assert 'BOVENLIGGEND (2)' in text
        assert 'D1  Betrouwbare beveiliging' in text
        assert 'D2  Informatiegestuurd werken' in text
        assert 'ONDERLIGGEND (1)' in text
        assert 'C1  Meldkamer' in text
        assert 'Digitalisering' in text
        assert 'Brigade Noord (eigenaar)' in text

    def test_element_zonder_verbindingen(self):
        text = all_text(self.slide_for(build({'slideColumns': ['Project']}), 'P2 — Los project'))
        assert 'Geen verbindingen.' in text
        assert 'Geen beschrijving ingevuld.' in text
        assert 'geen verbindingen met andere zichtbare elementen' in text

    def test_toelichting_bij_relaties_staat_er_niet_in(self):
        assert 'GEHEIM-TOELICHTING' not in full_text(build({'slideColumns': ['Benefit', 'Doel', 'Project']}))

    def test_grote_keten_wordt_per_kolom_begrensd(self):
        data = make_data()
        for i in range(10, 22):
            data['elements'].append(el(f'P{i}', 'Project'))
            data['edges'].append(edge(f'P{i}', 'C1'))
        slide = self.slide_for(build(self.OPTIONS, data), 'B1 — Eerste benefit')
        projects = [n for n in shape_names(slide) if n.startswith('element P')]
        assert len(projects) == CHAIN_MAX_ROWS - 1
        assert f'+ {13 - (CHAIN_MAX_ROWS - 1)} meer' in all_text(slide)

    def test_lange_beschrijving_loopt_door_op_vervolgslide(self):
        data = make_data()
        data['elements'][4]['description'] = 'Lange zin met inhoud. ' * 100
        prs = build(self.OPTIONS, data)
        assert len(prs.slides) == 6
        assert 'VERVOLG' in all_text(prs.slides[5])
        assert 'B1 — Eerste benefit' in all_text(prs.slides[5])

    def test_extreem_lange_tekst_geeft_hooguit_drie_slides_per_element(self):
        data = make_data()
        data['elements'][4]['description'] = 'x' * 200_000
        prs = build(self.OPTIONS, data)
        assert len(prs.slides) == 4 + 3

    def test_alias_type_hoort_bij_de_kolom_en_krijgt_eigen_kleur(self):
        data = make_data()
        data['columns'][2]['aliases'] = [{'typeName': 'Randvoorwaarde', 'color': '#111111'}]
        data['elements'].append(el('R1', 'Randvoorwaarde', 'Een randvoorwaarde'))
        prs = build(self.OPTIONS, data)
        slide = self.slide_for(prs, 'R1 — Een randvoorwaarde')
        assert 'BENEFIT (RANDVOORWAARDE)' in all_text(slide)
        node = next(s for s in slide.shapes if s.name == 'element R1 (geselecteerd)')
        assert str(node.fill.fore_color.rgb) == '111111'

    def test_paginanummers_lopen_door(self):
        prs = build(self.OPTIONS)
        for index, slide in enumerate(prs.slides, start=1):
            assert all_text(slide).rstrip().split('\n')[-1] == str(index)


class TestKlikbareVerbindingen:
    """DOEL-100: beschrijving links op maximaal 1/3, verbindingen rechts als
    blokjes in de kleur van hun kolom, klikbaar naar de eigen slide."""
    OPTIONS = {'slideColumns': ['Benefit', 'Doel']}

    def slide_for(self, prs, title):
        return next(s for s in prs.slides if title in all_text(s))

    def shape(self, slide, name):
        return next(s for s in slide.shapes if s.name == name)

    def test_beschrijving_maximaal_een_derde_en_blokjes_rechts(self):
        slide = self.slide_for(build(self.OPTIONS), 'B1 — Eerste benefit')
        desc = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith('BESCHRIJVING'))
        assert desc.width <= CONTENT_W / 3 + 1
        assert desc.left == MARGIN
        for chip in (s for s in slide.shapes if s.name.startswith('verbinding ')):
            assert chip.left >= RIGHT_LEFT
            assert chip.left + chip.width <= MARGIN + CONTENT_W + 1
        # Tags en organisatieonderdelen staan nu links bij de beschrijving.
        assert 'Digitalisering' in desc.text_frame.text and 'Brigade Noord' in desc.text_frame.text

    def test_blokje_in_kolomkleur_en_ondersteunend_gestippeld(self):
        slide = self.slide_for(build(self.OPTIONS), 'B1 — Eerste benefit')
        d1 = self.shape(slide, 'verbinding D1')
        d2 = self.shape(slide, 'verbinding D2 (ondersteunend)')
        c1 = self.shape(slide, 'verbinding C1')
        assert str(d1.line.color.rgb) == '8FAADC' and d1.line.dash_style is None
        assert str(d2.line.color.rgb) == '8FAADC' and d2.line.dash_style == MSO_LINE.DASH
        assert str(c1.line.color.rgb) == '6B4C8A'
        assert str(c1.fill.fore_color.rgb) == str(_tint(_rgb('#6B4C8A'), 0.8))

    def test_blokje_en_ketenvak_springen_naar_de_slide_van_het_element(self):
        prs = build(self.OPTIONS)
        b1 = self.slide_for(prs, 'B1 — Eerste benefit')
        d1_slide = next(s for s in prs.slides if 'D1 — Betrouwbare beveiliging' in all_text(s))
        for name in ('verbinding D1', 'element D1'):
            action = self.shape(b1, name).click_action
            assert action.action == PP_ACTION.NAMED_SLIDE
            assert action.target_slide.slide_id == d1_slide.slide_id
        # En terug: vanaf D1 naar B1 (een eerdere slide).
        back = self.shape(d1_slide, 'verbinding B1').click_action
        assert back.target_slide.slide_id == b1.slide_id

    def test_tooltip_noemt_het_doelelement(self):
        b1 = self.slide_for(build(self.OPTIONS), 'B1 — Eerste benefit')
        link = self.shape(b1, 'verbinding D2 (ondersteunend)')._element.find('.//' + qn('a:hlinkClick'))
        assert link.get('tooltip') == 'Ga naar D2 — Informatiegestuurd werken (ondersteunende relatie)'

    def test_geen_link_naar_element_zonder_eigen_slide(self):
        b1 = self.slide_for(build(self.OPTIONS), 'B1 — Eerste benefit')
        for name in ('verbinding C1', 'element C1', 'element M1', 'element B1 (geselecteerd)'):
            assert self.shape(b1, name).click_action.target_slide is None

    def test_legenda_noemt_klikken_alleen_als_er_links_zijn(self):
        with_links = self.slide_for(build(self.OPTIONS), 'B1 — Eerste benefit')
        assert 'Klik op een element' in all_text(with_links)
        without = self.slide_for(build({'slideColumns': ['Benefit']}), 'B1 — Eerste benefit')
        assert 'Klik op een element' not in all_text(without)

    def many_children(self, n):
        data = make_data()
        for i in range(n):
            data['elements'].append(el(f'X{i}', 'Capability', f'Capability met een redelijk lange naam {i}'))
            data['edges'].append(edge(f'X{i}', 'B1', 'ondersteunend' if i % 3 == 0 else 'primair'))
        return data

    def test_twaalf_verbindingen_passen_op_een_slide_zonder_vervolg(self):
        prs = build({'slideColumns': ['Benefit']}, self.many_children(11))  # + C1 = 12 onderliggend
        assert len(prs.slides) == 5
        slide = self.slide_for(prs, 'B1 — Eerste benefit')
        chips = [s for s in slide.shapes if s.name.startswith('verbinding ')]
        assert len(chips) == 2 + 12
        for chip in chips:
            assert chip.top + chip.height <= SLIDE_H - Inches(0.5)

    def test_heel_veel_verbindingen_lopen_door_en_worden_samengevat(self):
        n = 500
        prs = build({'slideColumns': ['Benefit']}, self.many_children(n))
        b1_slides = [s for s in prs.slides if 'B1 — Eerste benefit' in all_text(s)]
        assert len(b1_slides) == MAX_SLIDES_PER_ELEMENT
        assert 'VERVOLG' in all_text(b1_slides[1])
        shown = sum(1 for s in b1_slides for sh in s.shapes
                    if sh.name.startswith('verbinding ') and sh.name != 'verbinding + meer')
        more = next(sh for sh in b1_slides[-1].shapes if sh.name == 'verbinding + meer')
        hidden = int(re.search(r'\+ (\d+) meer', more.text_frame.text).group(1))
        assert shown + hidden == 2 + n + 1  # 2 boven, n + C1 onder
        for slide in b1_slides:
            for chip in slide.shapes:
                assert chip.top + chip.height <= SLIDE_H

    def test_vervolgslide_zonder_blokjes_gebruikt_de_hele_breedte(self):
        data = make_data()
        data['elements'][4]['description'] = 'Lange zin met inhoud. ' * 100
        prs = build(self.OPTIONS, data)
        vervolg = next(s for s in prs.slides if 'VERVOLG' in all_text(s) and 'B1 — Eerste' in all_text(s))
        body = next(s for s in vervolg.shapes if s.has_text_frame and 'BESCHRIJVING (VERVOLG)' in s.text_frame.text)
        assert body.width == CONTENT_W

    def test_owasp_a03_namen_in_blokjes_en_tooltips_blijven_letterlijk(self):
        data = make_data()
        evil = '"/><a:hlinkClick action="ppaction://program" r:id="x"/> <script>&'
        data['elements'][1]['name'] = evil  # D1
        prs = build(self.OPTIONS, data)
        b1 = self.slide_for(prs, 'B1 — Eerste benefit')
        chip = self.shape(b1, 'verbinding D1')
        assert evil[:20] in chip.text_frame.text
        link = chip._element.find('.//' + qn('a:hlinkClick'))
        assert link.get('action') == 'ppaction://hlinksldjump'
        assert evil in link.get('tooltip')
        # Er komt nergens een extra (geinjecteerde) actie bij.
        for slide in prs.slides:
            for h in slide._element.iter(qn('a:hlinkClick')):
                assert h.get('action') == 'ppaction://hlinksldjump'

    def test_owasp_geen_externe_links_in_het_pakket(self):
        data = make_data()
        data['elements'][1]['name'] = 'https://evil.example/phish'
        content = build_tree_pptx(data, self.OPTIONS, META)
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            for name in z.namelist():
                if name.endswith('.rels'):
                    rels = z.read(name).decode()
                    assert 'TargetMode="External"' not in rels
                    assert 'evil.example' not in rels


class TestTotaleBoom:
    """DOEL-103: slide 2 is de totale boom -- alle zichtbare kolommen met al
    hun elementen en relaties, op één slide, met klikbare kolomkoppen en
    vakjes."""
    OPTIONS = {'slideColumns': ['Benefit', 'Doel']}

    def overview(self, prs):
        return prs.slides[1]

    def shape(self, slide, name):
        return next(s for s in slide.shapes if s.name == name)

    def big(self, n, type_name='Project', prefix='Q'):
        data = make_data()
        for i in range(n):
            data['elements'].append(el(f'{prefix}{i}', type_name, f'Element met een redelijk lange naam {i}'))
            data['edges'].append(edge(f'{prefix}{i}', 'C1' if type_name == 'Project' else 'B1'))
        return data

    def test_slide_twee_is_de_totale_boom_met_alle_kolommen_en_elementen(self):
        prs = build(self.OPTIONS)
        slide = self.overview(prs)
        assert 'DE TOTALE BOOM' in all_text(slide) and 'Bewaken en Beveiligen' in all_text(slide)
        names = shape_names(slide)
        for code in ('M1', 'D1', 'D2', 'B1', 'B2', 'C1', 'C2', 'P1', 'P2'):
            assert f'overzicht-element {code}' in names
        assert len([n for n in names if n.startswith('overzicht-kolom ')]) == 5
        text = all_text(slide)
        for head in ('Project (2)', 'Capability (2)', 'Benefit (2)', 'Doel (2)', 'Missie (1)'):
            assert head in text
        assert 'B1 Eerste benefit' in text

    def test_alle_relaties_staan_erin_ook_de_ondersteunende(self):
        slide = self.overview(build(self.OPTIONS))
        lines = [n for n in shape_names(slide) if n.startswith('overzicht-relatie ')]
        assert len(lines) == len(make_data()['edges'])
        dashed = self.shape(slide, 'overzicht-relatie B1 > D2 (ondersteunend)')
        assert dashed.line.dash_style == MSO_LINE.DASH
        assert self.shape(slide, 'overzicht-relatie B1 > D1').line.dash_style is None
        assert 'stippellijn' in all_text(slide)

    def test_element_zonder_relaties_staat_er_ook_in(self):
        slide = self.overview(build(self.OPTIONS))
        assert 'overzicht-element P2' in shape_names(slide)
        assert not any('P2' in n for n in shape_names(slide) if n.startswith('overzicht-relatie '))

    def test_vak_in_kolomkleur_en_alias_in_eigen_kleur(self):
        data = make_data()
        data['columns'][2]['aliases'] = [{'typeName': 'Randvoorwaarde', 'color': '#111111'}]
        data['elements'].append(el('R1', 'Randvoorwaarde'))
        slide = self.overview(build(self.OPTIONS, data))
        assert str(self.shape(slide, 'overzicht-element B1').line.color.rgb) == 'C05A2C'
        assert str(self.shape(slide, 'overzicht-element R1').line.color.rgb) == '111111'
        assert str(self.shape(slide, 'overzicht-kolom 3').fill.fore_color.rgb) == 'C05A2C'

    def test_kolomkop_springt_naar_de_tussenslide_van_de_kolom(self):
        prs = build(self.OPTIONS)
        slide = self.overview(prs)
        sections = {  # kolomnummer -> tussenslide
            3: next(s for s in prs.slides if 'KOLOM' in all_text(s) and 'Benefit' in all_text(s)),
            4: next(s for s in prs.slides if 'KOLOM' in all_text(s) and 'Doel' in all_text(s)),
        }
        for number, target in sections.items():
            head = self.shape(slide, f'overzicht-kolom {number}')
            assert head.click_action.action == PP_ACTION.NAMED_SLIDE
            assert head.click_action.target_slide.slide_id == target.slide_id
        link = self.shape(slide, 'overzicht-kolom 3')._element.find('.//' + qn('a:hlinkClick'))
        assert link.get('tooltip') == 'Ga naar de kolom Benefit'

    def test_geen_link_bij_een_kolom_buiten_de_export(self):
        slide = self.overview(build(self.OPTIONS))
        for number in (1, 2, 5):  # Project, Capability, Missie zijn niet uitgewerkt
            assert self.shape(slide, f'overzicht-kolom {number}').click_action.target_slide is None
        for code in ('P1', 'C1', 'M1'):
            assert self.shape(slide, f'overzicht-element {code}').click_action.target_slide is None

    def test_vakje_springt_naar_de_slide_van_het_element(self):
        prs = build(self.OPTIONS)
        slide = self.overview(prs)
        for code, title in (('B1', 'B1 — Eerste benefit'), ('D2', 'D2 — Informatiegestuurd werken')):
            target = next(s for s in prs.slides if title in all_text(s))
            action = self.shape(slide, f'overzicht-element {code}').click_action
            assert action.action == PP_ACTION.NAMED_SLIDE
            assert action.target_slide.slide_id == target.slide_id
        link = self.shape(slide, 'overzicht-element B1')._element.find('.//' + qn('a:hlinkClick'))
        assert link.get('tooltip') == 'Ga naar B1 — Eerste benefit'

    def test_legenda_noemt_klikken_alleen_als_er_links_zijn(self):
        assert 'Klik op een kolomkop' in all_text(self.overview(build(self.OPTIONS)))
        assert 'Klik' not in all_text(self.overview(build({})))

    def test_verborgen_kolom_staat_niet_in_de_totale_boom_en_relatie_loopt_door(self):
        slide = self.overview(build({'visibleColumns': ['Project', 'Benefit', 'Doel', 'Missie']}))
        names = shape_names(slide)
        assert not any('C1' in n or 'C2' in n for n in names)
        assert len([n for n in names if n.startswith('overzicht-kolom ')]) == 4
        assert 'overzicht-relatie P1 > B1' in names

    @pytest.mark.parametrize('n', [34, 58, 200])
    def test_schaalt_bij_veel_elementen_en_blijft_binnen_de_slide(self, n):
        slide = self.overview(build({'slideColumns': ['Benefit']}, self.big(n)))
        nodes = [s for s in slide.shapes if s.name.startswith('overzicht-element ') or s.name.startswith('overzicht-meer ')]
        for node in nodes:
            assert node.top >= OVERVIEW_NODES_TOP - 1
            assert node.top + node.height <= OVERVIEW_NODES_BOTTOM + 1
            assert MARGIN <= node.left and node.left + node.width <= MARGIN + CONTENT_W + 1
        # Vakken in één kolom overlappen elkaar niet.
        projects = sorted((s for s in nodes if s.left == MARGIN), key=lambda s: s.top)
        for a, b in zip(projects, projects[1:]):
            assert a.top + a.height <= b.top

    def test_vierendertig_elementen_staan_er_allemaal_in(self):
        slide = self.overview(build({}, self.big(34)))
        projects = [n for n in shape_names(slide) if n.startswith('overzicht-element Q')]
        assert len(projects) == 34
        assert not any(n.startswith('overzicht-meer') for n in shape_names(slide))
        # Kleiner lettertype dan een kolom met weinig elementen.
        size = lambda name: self.shape(slide, name).text_frame.paragraphs[0].runs[0].font.size  # noqa: E731
        assert size('overzicht-element Q0') < size('overzicht-element M1')

    def test_owasp_a04_heel_grote_kolom_wordt_begrensd_met_plus_n_meer(self):
        slide = self.overview(build({}, self.big(500)))
        names = shape_names(slide)
        in_column = [n for n in names if n.startswith('overzicht-element Q') or n.startswith('overzicht-element P')]
        assert len(in_column) == OVERVIEW_MAX_ROWS - 1
        assert f'+ {502 - (OVERVIEW_MAX_ROWS - 1)} meer' in all_text(slide)
        assert 'Project (502)' in all_text(slide)

    def test_owasp_a04_aantal_lijnen_is_begrensd(self):
        data = {'columns': [col(0, 'A'), col(1, 'B')],
                'elements': [el(f'A{i}', 'A') for i in range(59)] + [el(f'B{i}', 'B') for i in range(59)],
                'edges': [edge(f'A{i}', f'B{j}', 'ondersteunend' if (i + j) % 2 else 'primair')
                          for i in range(59) for j in range(59)]}
        slide = self.overview(build({}, data))
        lines = [n for n in shape_names(slide) if n.startswith('overzicht-relatie ')]
        assert len(lines) == OVERVIEW_MAX_LINES
        assert f'Niet alle relaties zijn getekend ({OVERVIEW_MAX_LINES} van {59 * 59})' in all_text(slide)
        # Primaire relaties gaan voor.
        assert sum(1 for n in lines if not n.endswith('(ondersteunend)')) == (59 * 59 + 1) // 2

    def test_owasp_a03_namen_blijven_letterlijke_tekst(self):
        data = make_data()
        evil = '<a:hlinkClick r:id="rId9"/>"&\'><script>alert(1)</script>'
        data['elements'][4]['name'] = evil
        data['columns'][2]['title'] = 'Benefit "&<x>'
        content = build_tree_pptx(data, self.OPTIONS, META)
        slide = self.overview(load(content))
        assert '<script>' in self.shape(slide, 'overzicht-element B1').text_frame.text
        assert 'Benefit "&<x> (2)' in all_text(slide)
        link = self.shape(slide, 'overzicht-element B1')._element.find('.//' + qn('a:hlinkClick'))
        assert link.get('tooltip').startswith('Ga naar B1 — <a:hlinkClick')
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            xml = z.read('ppt/slides/slide2.xml').decode('utf-8')
        assert '<script>' not in xml and '&lt;script&gt;' in xml

    def test_owasp_alleen_interne_sprongen_in_het_pakket(self):
        content = build_tree_pptx(self.big(40), {'slideColumns': ['Project', 'Benefit', 'Doel', 'Missie', 'Capability']}, META)
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            for name in z.namelist():
                if not name.endswith('.rels'):
                    continue
                rels = z.read(name).decode('utf-8')
                assert 'TargetMode="External"' not in rels
                assert 'http://' not in re.sub(r'xmlns="[^"]+"|Type="[^"]+"', '', rels)
            slide2 = z.read('ppt/slides/slide2.xml').decode('utf-8')
        assert 'ppaction://hlinksldjump' in slide2
        assert not re.search(r'action="(?!ppaction://hlinksldjump)', slide2)

    def test_lege_boom_geeft_een_totale_boom_zonder_elementen(self):
        data = {'columns': [col(0, 'A'), col(1, 'B')], 'elements': [], 'edges': []}
        slide = self.overview(build({'slideColumns': ['A']}, data))
        assert 'A (0)' in all_text(slide) and 'B (0)' in all_text(slide)
        assert self.shape(slide, 'overzicht-kolom 1').click_action.target_slide is not None


class TestBoomfilter:
    """DOEL-107: met visibleElements bevat de hele presentatie alleen de
    gefilterde elementen; met filterPicks alleen de lijnen op het pad van de
    gekozen elementen, zoals op het scherm."""
    # Boomfilter op B1: B1 met alles erboven (D1, D2, M1) en eronder (C1, P1).
    KEEP = ['B1', 'D1', 'D2', 'M1', 'C1', 'P1']
    OPTIONS = {'visibleElements': KEEP, 'filterPicks': ['B1'], 'slideColumns': ['Benefit', 'Capability']}

    def shape(self, slide, name):
        return next(s for s in slide.shapes if s.name == name)

    def test_totale_boom_toont_alleen_de_gefilterde_elementen(self):
        slide = build(self.OPTIONS).slides[1]
        shown = sorted(n.removeprefix('overzicht-element ') for n in shape_names(slide) if n.startswith('overzicht-element '))
        assert shown == sorted(self.KEEP)
        text = all_text(slide)
        assert 'DE GEFILTERDE BOOM' in text and 'DE TOTALE BOOM' not in text
        assert 'Gefilterde boom: 6 van 9 elementen.' in text
        for head in ('Project (1)', 'Capability (1)', 'Benefit (1)', 'Doel (2)', 'Missie (1)'):
            assert head in text

    def test_weggefilterde_elementen_komen_nergens_in_de_presentatie_voor(self):
        prs = build(self.OPTIONS)
        text = full_text(prs)
        for gone in ('Tweede benefit', 'Sensoren', 'Los project'):
            assert gone not in text
        names = [n for s in prs.slides for n in shape_names(s)]
        assert not any(n.endswith(' B2') or n.endswith(' C2') or n.endswith(' P2') for n in names)

    def test_snoer_telt_de_gefilterde_aantallen_en_meldt_het_filter(self):
        text = all_text(build(self.OPTIONS).slides[0])
        assert 'Gefilterde boom: 6 van 9 elementen.' in text
        assert '2 elementen' in text and '1 element' in text  # Doel 2, de rest 1

    def test_kolomslides_en_elementslides_alleen_voor_gefilterde_elementen(self):
        prs = build(self.OPTIONS)
        # snoer + boom + (Capability: tussenslide + C1) + (Benefit: tussenslide + B1)
        assert len(prs.slides) == 2 + 2 + 2
        benefit = next(s for s in prs.slides if 'KOLOM' in all_text(s) and 'Benefit' in all_text(s))
        assert [n for n in shape_names(benefit) if n.startswith('kolomelement ')] == ['kolomelement B1']
        assert '1 element' in all_text(benefit)
        assert not any('B2 — Tweede benefit' in all_text(s) for s in prs.slides)

    def test_keten_en_verbindingen_op_de_elementslide_volgen_het_filter(self):
        data = make_data()
        data['edges'].append(edge('C2', 'B1'))  # C2 hangt ook onder B1, maar valt buiten het filter
        prs = build(self.OPTIONS, data)
        b1 = next(s for s in prs.slides if 'B1 — Eerste benefit' in all_text(s))
        names = shape_names(b1)
        assert 'verbinding C1' in names and 'verbinding C2' not in names
        assert 'element C2' not in names
        assert 'ONDERLIGGEND (1)' in all_text(b1)

    def test_zonder_picks_blijven_alle_relaties_tussen_zichtbare_elementen_staan(self):
        data = make_data()
        data['edges'].append(edge('P1', 'D1'))  # slaat B1 over
        slide = build({'visibleElements': self.KEEP}, data).slides[1]
        assert 'overzicht-relatie P1 > D1' in shape_names(slide)

    def test_met_picks_alleen_de_lijnen_op_het_pad_zoals_op_het_scherm(self):
        data = make_data()
        data['edges'].append(edge('P1', 'D1'))  # nakomeling rechtstreeks naar voorouder, om B1 heen
        slide = build(self.OPTIONS, data).slides[1]
        names = shape_names(slide)
        assert 'overzicht-relatie P1 > D1' not in names
        for kept in ('overzicht-relatie P1 > C1', 'overzicht-relatie C1 > B1', 'overzicht-relatie B1 > D1',
                     'overzicht-relatie B1 > D2 (ondersteunend)', 'overzicht-relatie D1 > M1', 'overzicht-relatie D2 > M1'):
            assert kept in names

    def test_gesnoeide_relatie_staat_ook_niet_bij_verbonden_met(self):
        data = make_data()
        data['edges'].append(edge('P1', 'D1'))
        prs = build({**self.OPTIONS, 'slideColumns': ['Project']}, data)
        p1 = next(s for s in prs.slides if 'P1 — Nieuw systeem' in all_text(s))
        assert 'verbinding C1' in shape_names(p1) and 'verbinding D1' not in shape_names(p1)

    def test_meerdere_picks_houden_de_paden_van_elk_van_hen(self):
        keep = ['B1', 'B2', 'D1', 'D2', 'M1', 'C1', 'C2', 'P1']
        slide = build({'visibleElements': keep, 'filterPicks': ['C1', 'C2']}).slides[1]
        names = shape_names(slide)
        for kept in ('overzicht-relatie C1 > B1', 'overzicht-relatie C2 > B2', 'overzicht-relatie P1 > C1',
                     'overzicht-relatie B2 > D2', 'overzicht-relatie D1 > M1'):
            assert kept in names

    def test_weggefilterd_element_wordt_niet_doorgetrokken_zoals_een_verborgen_kolom(self):
        # C1 valt buiten het filter: P1 mag dan niet "via C1" aan B1 hangen.
        slide = build({'visibleElements': ['P1', 'B1', 'D1', 'M1']}).slides[1]
        assert not any(n.startswith('overzicht-relatie P1') for n in shape_names(slide))

    def test_verborgen_kolom_binnen_het_filter_wordt_nog_wel_doorgetrokken(self):
        options = {'visibleElements': self.KEEP, 'filterPicks': ['B1'],
                   'visibleColumns': ['Project', 'Benefit', 'Doel', 'Missie']}
        slide = build(options).slides[1]
        assert 'overzicht-relatie P1 > B1' in shape_names(slide)
        assert 'Gefilterde boom: 5 van 7 elementen.' in all_text(slide)  # Capability telt niet mee

    def test_links_en_home_knop_werken_ook_in_de_gefilterde_presentatie(self):
        prs = build(self.OPTIONS)
        b1 = next(s for s in prs.slides if 'B1 — Eerste benefit' in all_text(s))
        assert self.shape(prs.slides[1], 'overzicht-element B1').click_action.target_slide.slide_id == b1.slide_id
        assert self.shape(prs.slides[1], 'overzicht-element D1').click_action.target_slide is None
        assert self.shape(b1, 'home-knop').click_action.target_slide.slide_id == prs.slides[1].slide_id

    def test_slidetelling_en_maximum_gelden_voor_de_gefilterde_set(self):
        data = {'columns': [col(0, 'A')], 'elements': [el(f'A{i}', 'A') for i in range(MAX_SLIDES + 50)], 'edges': []}
        with pytest.raises(TreePptxError):
            build_tree_pptx(data, {'slideColumns': ['A']}, META)
        prs = build({'slideColumns': ['A'], 'visibleElements': ['A1', 'A2', 'A3']}, data)
        assert len(prs.slides) == 2 + 1 + 3

    def test_zonder_filter_geen_melding_en_alle_elementen(self):
        prs = build({'slideColumns': ['Benefit']})
        assert 'Gefilterde boom' not in full_text(prs)
        assert 'DE TOTALE BOOM' in all_text(prs.slides[1])

    def test_leeg_filter_geeft_een_lege_boom_en_geen_fout(self):
        prs = build({'visibleElements': []})
        assert 'Gefilterde boom: 0 van 9 elementen.' in all_text(prs.slides[1])
        assert not any(n.startswith('overzicht-element ') for n in shape_names(prs.slides[1]))

    def test_onbekende_codes_en_verkeerde_types_worden_genegeerd(self):
        options = {'visibleElements': ['B1', 'BESTAAT-NIET', 5, None, {'x': 1}, 'B1'],
                   'filterPicks': ['BESTAAT-NIET', 7, 'B2']}
        slide = build(options).slides[1]
        assert [n for n in shape_names(slide) if n.startswith('overzicht-element ')] == ['overzicht-element B1']
        assert 'BESTAAT-NIET' not in all_text(slide)
        # Geen lijst = geen filter.
        assert 'Gefilterde boom' not in full_text(build({'visibleElements': 'B1', 'filterPicks': 'B1'}))

    def test_owasp_a04_cyclus_binnen_het_filter_loopt_niet_vast(self):
        data = make_data()
        data['edges'] += [edge('M1', 'P1'), edge('D1', 'B1')]
        prs = build(self.OPTIONS, data)
        assert len(prs.slides) == 6


class TestHomeKnop:
    """DOEL-106: elke slide behalve de totale boom heeft een Home-knop terug
    naar de totale boom (slide 2)."""
    OPTIONS = {'slideColumns': ['Benefit', 'Doel']}

    def home(self, slide):
        return [s for s in slide.shapes if s.name == 'home-knop']

    def check_all(self, prs, overview_index=1):
        overview = prs.slides[overview_index]
        assert 'DE TOTALE BOOM' in all_text(overview)
        for index, slide in enumerate(prs.slides):
            buttons = self.home(slide)
            if index == overview_index:
                assert buttons == []
                continue
            assert len(buttons) == 1, f'slide {index + 1}'
            action = buttons[0].click_action
            assert action.action == PP_ACTION.NAMED_SLIDE
            assert action.target_slide.slide_id == overview.slide_id

    def test_op_elke_slide_behalve_de_totale_boom_en_springt_naar_slide_twee(self):
        prs = build(self.OPTIONS)
        assert len(prs.slides) == 2 + (1 + 2) + (1 + 2)
        self.check_all(prs)

    def test_ook_zonder_uitgewerkte_kolommen_staat_hij_op_het_snoer(self):
        prs = build({})
        assert len(self.home(prs.slides[0])) == 1 and self.home(prs.slides[1]) == []
        self.check_all(prs)

    def test_ook_op_vervolgslides_van_een_element(self):
        data = make_data()
        data['elements'][4]['description'] = 'Lange zin met inhoud. ' * 400
        prs = build({'slideColumns': ['Benefit']}, data)
        assert any('VERVOLG' in all_text(s) for s in prs.slides)
        self.check_all(prs)

    def test_bij_meerdere_snoerslides_springt_hij_naar_de_totale_boom_erna(self):
        data = {'columns': [col(i, f'K{i}') for i in range(11)], 'elements': [el('A1', 'K0')], 'edges': []}
        prs = build({'perRow': 2, 'slideColumns': ['K0']}, data)  # 2 snoerslides, totale boom is slide 3
        self.check_all(prs, overview_index=2)

    def test_knop_met_tekst_en_tooltip(self):
        button = self.home(build(self.OPTIONS).slides[0])[0]
        assert button.text_frame.text == 'Home'
        assert button._element.find('.//' + qn('a:hlinkClick')).get('tooltip') == 'Terug naar de totale boom'

    def test_vaste_plek_in_de_voettekst_naast_het_paginanummer(self):
        prs = build(self.OPTIONS)
        positions = set()
        for index, slide in enumerate(prs.slides):
            if index == 1:
                continue
            button = self.home(slide)[0]
            positions.add((button.left, button.top, button.width, button.height))
            assert button.top + button.height <= SLIDE_H
            assert button.top >= SLIDE_H - Inches(0.5)  # onder de inhoud (die stopt op 0,5 inch van de rand)
            number = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text == str(index + 1)][-1]
            assert button.left + button.width <= number.left
            footer = next(s for s in slide.shapes if s.has_text_frame and 'Bewaken en Beveiligen — CLPK' in s.text_frame.text)
            assert footer.left + footer.width <= button.left
        assert len(positions) == 1

    def test_paginanummer_blijft_het_laatste_op_de_slide(self):
        prs = build(self.OPTIONS)
        for index, slide in enumerate(prs.slides, start=1):
            assert all_text(slide).rstrip().split('\n')[-1] == str(index)

    def test_owasp_home_knop_is_alleen_een_interne_sprong(self):
        content = build_tree_pptx(make_data(), self.OPTIONS, META)
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            for name in z.namelist():
                if name.endswith('.rels'):
                    assert 'TargetMode="External"' not in z.read(name).decode('utf-8')
            slide1 = z.read('ppt/slides/slide1.xml').decode('utf-8')
        assert re.search(r'name="home-knop"><a:hlinkClick action="ppaction://hlinksldjump"', slide1)


class TestKolomslideBlokjes:
    """DOEL-103: de tussenslide van een kolom toont alle elementen van die
    kolom als klikbare blokjes in de kolomkleur."""
    OPTIONS = {'slideColumns': ['Benefit', 'Doel']}

    def section(self, prs, title):
        return next(s for s in prs.slides if 'KOLOM' in all_text(s) and title in all_text(s))

    def chips(self, slide):
        return [s for s in slide.shapes if s.name.startswith('kolomelement ')]

    def many(self, n):
        data = make_data()
        data['elements'] += [el(f'X{i}', 'Benefit', f'Benefit met een redelijk lange naam {i}') for i in range(n)]
        return data

    def test_alle_elementen_van_de_kolom_als_blokje_in_kolomkleur(self):
        slide = self.section(build(self.OPTIONS), 'Benefit')
        chips = self.chips(slide)
        assert [c.name for c in chips] == ['kolomelement B2', 'kolomelement B1']  # boomvolgorde
        for chip in chips:
            assert str(chip.line.color.rgb) == 'C05A2C'
            assert str(chip.fill.fore_color.rgb) == str(_tint(_rgb('#C05A2C'), 0.8))
        assert 'B1  Eerste benefit' in all_text(slide)
        assert '2 elementen' in all_text(slide)

    def test_blokjes_staan_rechtsonder_onder_het_gekleurde_blok(self):
        slide = self.section(build(self.OPTIONS), 'Benefit')
        count = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == '2 elementen')
        for chip in self.chips(slide):
            assert chip.left >= count.left + count.width
            assert chip.top >= count.top
            assert chip.left + chip.width <= MARGIN + CONTENT_W + 1

    def test_blokje_springt_naar_de_slide_van_het_element(self):
        prs = build(self.OPTIONS)
        slide = self.section(prs, 'Benefit')
        target = next(s for s in prs.slides if 'B1 — Eerste benefit' in all_text(s))
        chip = next(c for c in self.chips(slide) if c.name == 'kolomelement B1')
        assert chip.click_action.action == PP_ACTION.NAMED_SLIDE
        assert chip.click_action.target_slide.slide_id == target.slide_id
        assert chip._element.find('.//' + qn('a:hlinkClick')).get('tooltip') == 'Ga naar B1 — Eerste benefit'

    def test_alias_element_krijgt_de_aliaskleur(self):
        data = make_data()
        data['columns'][2]['aliases'] = [{'typeName': 'Randvoorwaarde', 'color': '#111111'}]
        data['elements'].append(el('R1', 'Randvoorwaarde'))
        slide = self.section(build(self.OPTIONS, data), 'Benefit')
        chip = next(c for c in self.chips(slide) if c.name == 'kolomelement R1')
        assert str(chip.line.color.rgb) == '111111'

    @pytest.mark.parametrize('n', [1, 6, 32, 80, 130])
    def test_blokjes_schalen_mee_en_blijven_binnen_de_slide(self, n):
        prs = build({'slideColumns': ['Benefit']}, self.many(n))
        slide = self.section(prs, 'Benefit')
        chips = self.chips(slide)
        assert len(chips) == n + 2
        assert 'kolomelement + meer' not in shape_names(slide)
        band = max((s for s in slide.shapes if s.width > CONTENT_W and s.height < SLIDE_H), key=lambda s: s.height)
        for chip in chips:
            assert chip.top >= band.top + band.height
            assert chip.top + chip.height <= SLIDE_H - Inches(0.5)
            assert chip.left + chip.width <= MARGIN + CONTENT_W + 1
            assert chip.click_action.target_slide is not None

    def test_veel_elementen_geeft_kleinere_blokjes(self):
        few = self.chips(self.section(build({'slideColumns': ['Benefit']}, self.many(2)), 'Benefit'))
        lots = self.chips(self.section(build({'slideColumns': ['Benefit']}, self.many(120)), 'Benefit'))
        assert lots[0].width < few[0].width and lots[0].height < few[0].height

    def test_past_het_echt_niet_dan_plus_n_meer(self):
        slide = self.section(build({'slideColumns': ['Benefit']}, self.many(250)), 'Benefit')
        chips = [c for c in self.chips(slide) if c.name != 'kolomelement + meer']
        assert len(chips) == SECTION_MAX_CHIPS - 1
        more = next(s for s in slide.shapes if s.name == 'kolomelement + meer')
        assert more.text_frame.text == f'+ {252 - (SECTION_MAX_CHIPS - 1)} meer'
        assert more.click_action.target_slide is None
        assert more.top + more.height <= SLIDE_H - Inches(0.5)
        assert '252 elementen' in all_text(slide)

    def test_lege_kolom_heeft_geen_blokjes_en_geen_klikhint(self):
        data = make_data()
        data['columns'].append(col(5, 'Leeg'))
        slide = self.section(build({'slideColumns': ['Leeg']}, data), 'Leeg')
        assert self.chips(slide) == []
        assert '0 elementen' in all_text(slide) and 'Klik op een element' not in all_text(slide)

    def test_owasp_a03_naam_in_blokje_blijft_letterlijke_tekst(self):
        data = make_data()
        data['elements'][4]['name'] = '<script>alert(1)</script>'
        slide = self.section(build(self.OPTIONS, data), 'Benefit')
        chip = next(c for c in self.chips(slide) if c.name == 'kolomelement B1')
        assert '<script>' in chip.text_frame.text


class TestRobuustheid:
    """Vijandige of kapotte invoer mag de service niet laten crashen en geen
    ongeldig bestand opleveren (OWASP A03/A04)."""

    def test_cyclus_in_relaties_loopt_niet_vast(self):
        data = make_data()
        data['edges'] += [edge('M1', 'P1'), edge('D1', 'B1'), edge('B1', 'B1')]
        prs = build({'slideColumns': ['Benefit', 'Project', 'Missie']}, data)
        assert len(prs.slides) > 4

    def test_cyclus_door_verborgen_kolom_loopt_niet_vast(self):
        data = make_data()
        data['elements'].append(el('C3', 'Capability'))
        data['edges'] += [edge('C1', 'C3'), edge('C3', 'C1'), edge('C3', 'B2', 'ondersteunend')]
        model = _Model(data, {'visibleColumns': ['Project', 'Benefit']})
        assert model.parents['P1'] == {'B1': True, 'B2': False}

    def test_hele_diepe_keten_geeft_geen_recursiefout(self):
        n = 3000
        data = {
            'columns': [col(0, 'A'), col(1, 'Verborgen'), col(2, 'Z')],
            'elements': [el('A1', 'A'), el('Z1', 'Z')] + [el(f'H{i}', 'Verborgen') for i in range(n)],
            'edges': [edge('A1', 'H0'), edge(f'H{n - 1}', 'Z1')] + [edge(f'H{i}', f'H{i + 1}') for i in range(n - 1)],
        }
        model = _Model(data, {'visibleColumns': ['A', 'Z']})
        assert model.parents['A1'] == {'Z1': True}

    def test_xml_en_besturingstekens_in_tekst(self):
        data = make_data()
        data['elements'][4].update(
            name='<script>alert(1)</script> & "quotes"', description='Nul\x00byte \x0b en ]]> <a:t>inject</a:t>')
        data['columns'][2]['title'] = 'Benefit</a:t><evil/>'
        prs = build({'slideColumns': ['Benefit']}, data)
        text = full_text(prs)
        assert '<script>alert(1)</script> & "quotes"' in text
        assert 'Nulbyte' in text and '<a:t>inject</a:t>' in text
        assert '\x00' not in text

    def test_formule_achtige_tekst_blijft_letterlijke_tekst(self):
        data = make_data()
        data['elements'][4]['name'] = '=HYPERLINK("http://evil","klik")'
        prs = build({'slideColumns': ['Benefit']}, data)
        assert '=HYPERLINK("http://evil","klik")' in full_text(prs)
        for slide in prs.slides:
            for shape in slide.shapes:
                # Alleen interne sprongen naar een slide (DOEL-100), nooit een
                # (externe) hyperlink-adres uit de invoer.
                address = shape.click_action.hyperlink.address
                assert not address or re.fullmatch(r'slide\d+\.xml', address)
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        for run in p.runs:
                            assert run.hyperlink.address is None

    def test_verkeerde_types_in_de_invoer(self):
        data = {
            'columns': [col(0, 'A'), 'geen dict', {'typeName': None}, {'typeName': 'B', 'color': 5, 'aliases': 'x', 'position': 'nee'}],
            'elements': [el('A1', 'A'), None, 7, {'code': None}, {'code': 'A1', 'type': 'A'}, {'code': 'X', 'type': ['A']}],
            'edges': [None, 'x', {'source': 'A1'}, {'source': 'A1', 'target': 'BESTAAT-NIET'}, {'source': 1, 'target': 2}],
            'tags': 'geen lijst', 'elementTags': ['geen dict'], 'orgUnits': None, 'obOrg': {'A1': 'geen lijst'},
        }
        prs = load(build_tree_pptx(data, {'slideColumns': ['A', 5, None], 'visibleColumns': 'alles', 'perRow': []}, 'geen meta'))
        assert len(prs.slides) == 4

    def test_lege_of_ontbrekende_invoer(self):
        for data in (None, {}, [], 'tekst', {'columns': []}):
            with pytest.raises(TreePptxError):
                build_tree_pptx(data, None, None)

    def test_te_veel_slides_wordt_geweigerd_met_duidelijke_melding(self):
        data = {'columns': [col(0, 'A')], 'elements': [el(f'A{i}', 'A') for i in range(MAX_SLIDES)], 'edges': []}
        with pytest.raises(TreePptxError) as exc:
            build_tree_pptx(data, {'slideColumns': ['A']}, META)
        assert str(MAX_SLIDES) in str(exc.value)
        # Zonder element-slides mag dezelfde grote boom wel.
        assert len(build({}, data).slides) == 2

    def test_precies_op_de_grens_mag(self):
        # snoer + totale boom + tussenslide + een slide per element.
        data = {'columns': [col(0, 'A')], 'elements': [el(f'A{i}', 'A') for i in range(MAX_SLIDES - 3)], 'edges': []}
        assert len(build({'slideColumns': ['A']}, data).slides) == MAX_SLIDES

    def test_een_element_boven_de_grens_wordt_geweigerd(self):
        # DOEL-103: de totale boom telt mee, dus MAX_SLIDES - 2 elementen is er nu een te veel.
        data = {'columns': [col(0, 'A')], 'elements': [el(f'A{i}', 'A') for i in range(MAX_SLIDES - 2)], 'edges': []}
        with pytest.raises(TreePptxError):
            build_tree_pptx(data, {'slideColumns': ['A']}, META)


class TestEndpoint:
    def test_geeft_een_pptx_terug(self):
        res = client.post('/tree-pptx', json={'data': make_data(), 'options': {'slideColumns': ['Doel']}, 'meta': META})
        assert res.status_code == 200
        assert res.headers['content-type'] == (
            'application/vnd.openxmlformats-officedocument.presentationml.presentation')
        assert res.content[:2] == b'PK'
        assert len(load(res.content).slides) == 5

    def test_ongeldige_keuze_geeft_422_met_melding(self):
        res = client.post('/tree-pptx', json={'data': make_data(), 'options': {'visibleColumns': []}, 'meta': META})
        assert res.status_code == 422
        assert 'zichtbare kolom' in res.json()['error']

    def test_lege_body_geeft_422_en_geen_500(self):
        res = client.post('/tree-pptx', json={})
        assert res.status_code == 422
        assert 'error' in res.json()

    def test_foutmelding_lekt_geen_interne_details(self):
        res = client.post('/tree-pptx', json={'data': {'columns': 'kapot'}})
        assert res.status_code == 422
        assert 'Traceback' not in res.text and 'tree_pptx.py' not in res.text


class TestBroertjesNavigatie:
    """DOEL-117: op elke element-slide (ook vervolgslides) springen Vorige en
    Volgende naar het vorige/volgende broertje: zelfde kolom, zelfde primaire
    ouder, boomvolgorde. Zonder broertje grijs en zonder link."""
    OPTIONS = {'slideColumns': ['Capability']}

    def data(self):
        data = make_data()
        # C1 en C3 hangen primair onder B1, C2 alleen onder B2.
        data['elements'].insert(6, el('C3', 'Capability', 'Camera\'s'))
        data['edges'].append(edge('C3', 'B1'))
        return data

    def slide_of(self, prs, code):
        return next(s for s in prs.slides if f'{code} — ' in all_text(s) and 'VERVOLG' not in all_text(s))

    def button(self, slide, name):
        found = [s for s in slide.shapes if s.name == name]
        assert len(found) == 1, f'{name}: {len(found)} keer'
        return found[0]

    def target(self, slide, name):
        b = self.button(slide, name)
        if b._element.find('.//' + qn('a:hlinkClick')) is None:
            return None
        assert b.click_action.action == PP_ACTION.NAMED_SLIDE
        return b.click_action.target_slide

    def assert_nav(self, prs, code, prev, nxt):
        slide = self.slide_of(prs, code)
        for name, want in (('vorige-knop', prev), ('volgende-knop', nxt)):
            got = self.target(slide, name)
            if want is None:
                assert got is None, f'{code} {name}: verwacht grijs'
                assert str(self.button(slide, name).fill.fore_color.rgb) == 'E1E4E9'
            else:
                assert got is not None, f'{code} {name}: verwacht link naar {want}'
                assert got.slide_id == self.slide_of(prs, want).slide_id, f'{code} {name}'

    def test_broertjes_onder_dezelfde_ouder_in_boomvolgorde(self):
        prs = build(self.OPTIONS, self.data())
        self.assert_nav(prs, 'C1', None, 'C3')
        self.assert_nav(prs, 'C3', 'C1', None)
        self.assert_nav(prs, 'C2', None, None)  # enig kind van B2

    def test_volgorde_is_de_boomvolgorde(self):
        data = self.data()
        c3 = data['elements'].pop(6)
        data['elements'].insert(5, c3)  # C3 vóór C1
        prs = build(self.OPTIONS, data)
        self.assert_nav(prs, 'C3', None, 'C1')
        self.assert_nav(prs, 'C1', 'C3', None)

    def test_knoppen_met_tekst_en_tooltip(self):
        slide = self.slide_of(build(self.OPTIONS, self.data()), 'C1')
        assert self.button(slide, 'vorige-knop').text_frame.text == '◄ Vorige'
        nxt = self.button(slide, 'volgende-knop')
        assert nxt.text_frame.text == 'Volgende ►'
        assert nxt._element.find('.//' + qn('a:hlinkClick')).get('tooltip') == "Volgende onder B1: C3 — Camera's"

    def test_vervolgslides_worden_overgeslagen_en_hebben_zelf_ook_de_knoppen(self):
        data = self.data()
        next(e for e in data['elements'] if e['code'] == 'C1')['description'] = 'Lange zin met inhoud. ' * 400
        prs = build(self.OPTIONS, data)
        c1_slides = [s for s in prs.slides if 'C1 — ' in all_text(s)]
        assert len(c1_slides) >= 2 and 'VERVOLG' in all_text(c1_slides[1])
        # C3 springt terug naar de EERSTE slide van C1, niet naar een vervolgslide.
        assert self.target(self.slide_of(prs, 'C3'), 'vorige-knop').slide_id == c1_slides[0].slide_id
        for s in c1_slides:
            assert self.target(s, 'volgende-knop').slide_id == self.slide_of(prs, 'C3').slide_id
            assert self.target(s, 'vorige-knop') is None

    def test_ondersteunende_relatie_telt_niet_als_ouder(self):
        data = self.data()
        data['elements'].append(el('C4', 'Capability', 'Alleen ondersteunend'))
        data['edges'].append(edge('C4', 'B1', 'ondersteunend'))
        prs = build(self.OPTIONS, data)
        self.assert_nav(prs, 'C3', 'C1', None)  # C4 is geen broertje
        self.assert_nav(prs, 'C4', None, None)  # geen primaire ouder

    def test_bij_meer_primaire_ouders_telt_de_eerste_in_boomvolgorde(self):
        data = self.data()
        data['edges'].append(edge('C3', 'B2'))  # B2 staat in de boom vóór B1
        prs = build(self.OPTIONS, data)
        self.assert_nav(prs, 'C1', None, None)
        self.assert_nav(prs, 'C3', None, 'C2')
        self.assert_nav(prs, 'C2', 'C3', None)

    def test_alias_in_dezelfde_kolom_is_ook_een_broertje(self):
        data = self.data()
        data['columns'][1]['aliases'] = [{'typeName': 'Besluit', 'color': '#C0392B'}]
        data['elements'].append(el('X1', 'Besluit', 'Een besluit'))
        data['edges'].append(edge('X1', 'B1'))
        prs = build(self.OPTIONS, data)
        self.assert_nav(prs, 'C3', 'C1', 'X1')
        self.assert_nav(prs, 'X1', 'C3', None)

    def test_boomfilter_broertje_buiten_het_filter_wordt_overgeslagen(self):
        data = self.data()
        keep = [e['code'] for e in data['elements'] if e['code'] != 'C3']
        prs = build({**self.OPTIONS, 'visibleElements': keep}, data)
        self.assert_nav(prs, 'C1', None, None)

    def test_verborgen_ouderkolom_dan_telt_de_eerstvolgende_zichtbare_ouder(self):
        data = self.data()
        data['edges'] = [e for e in data['edges'] if not (e['source'] == 'B2' and e['target'] == 'D2')]
        data['edges'].append(edge('B2', 'D1'))
        visible = ['Project', 'Capability', 'Doel', 'Missie']  # Benefit verborgen
        prs = build({**self.OPTIONS, 'visibleColumns': visible}, data)
        # Alle drie hangen nu (doorgetrokken) onder D1.
        self.assert_nav(prs, 'C1', None, 'C3')
        self.assert_nav(prs, 'C3', 'C1', 'C2')
        self.assert_nav(prs, 'C2', 'C3', None)

    def test_werkt_in_elke_uitgewerkte_kolom(self):
        # Benefit uitgewerkt: B2 en B1 worden broertjes zodra ze allebei primair onder D1 hangen.
        data = self.data()
        data['edges'].append(edge('B2', 'D1'))
        prs = build({'slideColumns': ['Benefit']}, data)
        self.assert_nav(prs, 'B2', None, 'B1')
        self.assert_nav(prs, 'B1', 'B2', None)

    def test_niet_op_snoer_totale_boom_en_kolomslides(self):
        prs = build(self.OPTIONS, self.data())
        for slide in prs.slides:
            names = shape_names(slide)
            is_element = any(re.search(r'\bC\d — ', line) for line in all_text(slide).split('\n')[:3])
            has = 'vorige-knop' in names or 'volgende-knop' in names
            assert has == is_element, all_text(slide)[:80]

    def test_plek_in_de_voettekst_links_van_home_zonder_overlap(self):
        prs = build(self.OPTIONS, self.data())
        positions = set()
        for index, slide in enumerate(prs.slides, start=1):
            names = shape_names(slide)
            if 'vorige-knop' not in names:
                continue
            prev, nxt = self.button(slide, 'vorige-knop'), self.button(slide, 'volgende-knop')
            home = self.button(slide, 'home-knop')
            positions.add((prev.left, nxt.left, home.left))
            assert prev.left + prev.width <= nxt.left and nxt.left + nxt.width <= home.left
            assert prev.top == nxt.top == home.top and prev.top + prev.height <= SLIDE_H
            footer = next(s for s in slide.shapes if s.has_text_frame and 'Bewaken en Beveiligen — CLPK' in s.text_frame.text)
            assert footer.left + footer.width <= prev.left
            assert all_text(slide).rstrip().split('\n')[-1] == str(index)
        assert len(positions) == 1

    def test_owasp_alleen_interne_sprongen_en_veilige_tooltip(self):
        data = self.data()
        next(e for e in data['elements'] if e['code'] == 'C3')['name'] = '"/><script>alert(1)</script>'
        content = build_tree_pptx(data, self.OPTIONS, META)
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            for name in z.namelist():
                if name.endswith('.rels'):
                    assert 'TargetMode="External"' not in z.read(name).decode('utf-8')
                if name.startswith('ppt/slides/slide') and name.endswith('.xml'):
                    xml = z.read(name).decode('utf-8')
                    assert '<script>' not in xml
                    for m in re.finditer(r'name="(vorige|volgende)-knop"><a:hlinkClick ([^>]*)', xml):
                        assert 'ppaction://hlinksldjump' in m.group(2)
        slide = self.slide_of(load(content), 'C1')
        tooltip = self.button(slide, 'volgende-knop')._element.find('.//' + qn('a:hlinkClick')).get('tooltip')
        assert tooltip == 'Volgende onder B1: C3 — "/><script>alert(1)</script>'

    def test_model_hulpfuncties(self):
        model = _Model(self.data(), self.OPTIONS)
        assert model.main_parent('C1') == 'B1'
        assert model.main_parent('M1') is None
        nav = model.sibling_nav('C3')
        assert (nav.prev, nav.next, nav.parent) == ('C1', None, 'B1')


class TestBroertjesInDeKeten:
    """DOEL-117: in de keten boven het geselecteerde element een schaduwkaartje
    van het vorige broertje (▲), eronder van het volgende (▼), klikbaar; de
    positie ("n van m") staat in de legenda, niet op het element."""
    OPTIONS = {'slideColumns': ['Capability']}

    def data(self):
        data = make_data()
        data['elements'].insert(5, el('C0', 'Capability', 'Uitgifte'))
        data['elements'].insert(7, el('C3', 'Capability', "Camera's"))
        data['edges'] += [edge('C0', 'B1'), edge('C3', 'B1')]  # onder B1: C0, C1, C3
        return data

    def slide_of(self, prs, code):
        return next(s for s in prs.slides if f'{code} — ' in all_text(s) and 'VERVOLG' not in all_text(s))

    def cards(self, slide):
        return {s.name: s for s in slide.shapes if s.name.startswith('broertje ')}

    def teller(self, slide):
        m = re.search(r'▲/▼ = vorig/volgend element onder (\S+) \((\d+ van \d+)\)\.', all_text(slide))
        return m.group(2) if m else None

    def test_geen_teller_meer_op_het_element_zelf(self):
        slide = self.slide_of(build(self.OPTIONS, self.data()), 'C1')
        assert not [s for s in slide.shapes if s.name == 'broertjes-teller']
        assert '▲/▼ = vorig/volgend element onder B1 (2 van 3).' in all_text(slide)

    def test_middelste_broertje_heeft_beide_kaartjes_en_teller(self):
        prs = build(self.OPTIONS, self.data())
        slide = self.slide_of(prs, 'C1')
        cards = self.cards(slide)
        assert set(cards) == {'broertje vorige C0', 'broertje volgende C3'}
        assert cards['broertje vorige C0'].text_frame.text == '▲ C0 Uitgifte'
        assert cards['broertje volgende C3'].text_frame.text == "▼ C3 Camera's"
        assert cards['broertje vorige C0'].click_action.target_slide.slide_id == self.slide_of(prs, 'C0').slide_id
        assert cards['broertje volgende C3'].click_action.target_slide.slide_id == self.slide_of(prs, 'C3').slide_id
        tip = cards['broertje volgende C3']._element.find('.//' + qn('a:hlinkClick')).get('tooltip')
        assert tip == "Volgende onder B1: C3 — Camera's"
        assert self.teller(slide) == '2 van 3'

    def test_eerste_en_laatste_hebben_alleen_het_kaartje_aan_de_kant_met_een_broertje(self):
        prs = build(self.OPTIONS, self.data())
        assert set(self.cards(self.slide_of(prs, 'C0'))) == {'broertje volgende C1'}
        assert self.teller(self.slide_of(prs, 'C0')) == '1 van 3'
        assert set(self.cards(self.slide_of(prs, 'C3'))) == {'broertje vorige C1'}
        assert self.teller(self.slide_of(prs, 'C3')) == '3 van 3'

    def test_enig_kind_geen_kaartjes_geen_teller_geen_legenda(self):
        slide = self.slide_of(build(self.OPTIONS, self.data()), 'C2')
        assert self.cards(slide) == {} and self.teller(slide) is None
        assert '▲/▼' not in all_text(slide)

    def test_kaartjes_boven_en_onder_het_element_binnen_het_ketenvlak(self):
        slide = self.slide_of(build(self.OPTIONS, self.data()), 'C1')
        node = next(s for s in slide.shapes if s.name == 'element C1 (geselecteerd)')
        cards = self.cards(slide)
        above, below = cards['broertje vorige C0'], cards['broertje volgende C3']
        assert above.top + above.height < node.top and below.top > node.top + node.height
        for c in (above, below):
            assert node.left < c.left and c.left + c.width < node.left + node.width  # smaller dan het vak
            assert c.top >= Inches(1.79) and c.top + c.height <= Inches(1.84) + Inches(2.36)

    def test_geen_kaartjes_als_er_meer_vakken_in_de_kolom_van_het_element_staan(self):
        data = self.data()
        data['edges'].append(edge('C1', 'C3'))  # relatie binnen de kolom: C3 staat dan in de keten van C1
        slide = self.slide_of(build(self.OPTIONS, data), 'C1')
        assert self.cards(slide) == {} and self.teller(slide) is None

    def test_niet_op_vervolgslides(self):
        data = self.data()
        next(e for e in data['elements'] if e['code'] == 'C1')['description'] = 'Lange zin met inhoud. ' * 400
        prs = build(self.OPTIONS, data)
        vervolg = [s for s in prs.slides if 'C1 — ' in all_text(s) and 'VERVOLG' in all_text(s)]
        assert vervolg and all(self.cards(s) == {} for s in vervolg)

    def test_owasp_kaartjes_zijn_alleen_interne_sprongen_met_veilige_tekst(self):
        data = self.data()
        next(e for e in data['elements'] if e['code'] == 'C3')['name'] = '<img src=x onerror=alert(1)>'
        content = build_tree_pptx(data, self.OPTIONS, META)
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            xml = ''.join(z.read(n).decode('utf-8') for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml'))
            rels = ''.join(z.read(n).decode('utf-8') for n in z.namelist() if n.endswith('.rels'))
        assert 'TargetMode="External"' not in rels
        assert '<img' not in xml
        for m in re.finditer(r'name="broertje [^"]*"><a:hlinkClick ([^>]*)', xml):
            assert 'ppaction://hlinksldjump' in m.group(1)
        assert re.search(r'name="broertje volgende C3"><a:hlinkClick', xml)
