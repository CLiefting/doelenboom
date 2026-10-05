"""
Export van een HELE doelenboom als PowerPoint-presentatie (DOEL-88).

Slide-structuur:
  1.   De kolommen als "snoer": elke zichtbare kolom een vak in de eigen
       kleur met titel en omschrijving, pijlen met het relatielabel ertussen,
       maximaal `perRow` vakken per regel en daarna slingerend verder op de
       volgende regel(s). Kolommen waarvoor element-slides volgen zijn
       gemarkeerd. Meer dan SNOER_MAX_ROWS regels loopt door op een
       vervolgslide.
  Per gekozen kolom:
  n.   Een tussenslide met de kolom.
  n+1. Eén slide per element, in boomvolgorde: bovenaan de boom gefilterd op
       dat ene element (het hele pad omhoog en omlaag, getekend met
       PowerPoint-vormen), daaronder de beschrijving, KPI, de directe
       verbindingen, tags en organisatieonderdelen. Lange tekst loopt door op
       een vervolgslide.

Verborgen kolommen (options.visibleColumns) komen nergens in de presentatie
voor: niet in het snoer, niet in de keten en niet onder "Verbonden met".
Relaties die via een verborgen kolom lopen worden doorgetrokken naar de
eerstvolgende zichtbare kolom (zie _visible_edges).

Puur een export, geen import. Bewust NIET in de presentatie: kenmerken,
motivaties van afwijkingen en projectgegevens (status, producten,
activiteiten) -- de API (routes/exports.ts) stuurt die ook niet mee.

Aangeroepen door api/src/routes/exports.ts (POST .../export-pptx) via
POST /tree-pptx in main.py.
"""
from __future__ import annotations

import io
import math
import re
import textwrap
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .project_pptx import (
    ACCENT,
    CONTENT_W,
    DARK,
    LIGHT_BG,
    MARGIN,
    MUTED,
    SLIDE_H,
    SLIDE_W,
    WHITE,
    _add_rect,
    _add_text,
    _blank_slide,
    _fmt_date,
    _new_presentation,
    _truncate,
)

# Bovengrens op het aantal slides (snoer + tussenslides + element-slides,
# vóór eventuele vervolgslides). Dezelfde grens staat in de API
# (routes/exports.ts) en in het dialoogvenster (tree.html).
MAX_SLIDES = 300
PER_ROW_MIN = 2
PER_ROW_MAX = 6
PER_ROW_DEFAULT = 4
SNOER_MAX_ROWS = 4
# Maximaal aantal vakken per kolom in de getekende keten; de rest wordt
# samengevat als "+ n meer".
CHAIN_MAX_ROWS = 6
# Een element krijgt hooguit zoveel slides (1 + vervolg); daarna wordt de
# tekst ingekort.
MAX_SLIDES_PER_ELEMENT = 3

LINE_COLOR = RGBColor(0x8A, 0x8F, 0x98)
FALLBACK_COLOR = RGBColor(0x6C, 0x6F, 0x76)
_HEX_RE = re.compile(r'^#[0-9a-fA-F]{6}$')
# Besturingstekens die niet in XML mogen staan (tab/newline blijven).
_CONTROL_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]')


class TreePptxError(ValueError):
    """Ongeldige invoer of keuze (bv. te veel slides) -- main.py maakt hier
    een 422 van met deze melding."""


# ---- kleine hulpfuncties -------------------------------------------------

def _clean(value: Any) -> str:
    if value is None:
        return ''
    return _CONTROL_RE.sub('', str(value)).strip()


def _rgb(hex_color: Any) -> RGBColor:
    """'#3E6FA6' -> RGBColor; alles wat geen geldige hex-kleur is valt terug
    op grijs (de kleur komt uit de kolomconfiguratie, maar wordt hier niet
    blind vertrouwd)."""
    if isinstance(hex_color, str) and _HEX_RE.match(hex_color):
        return RGBColor.from_string(hex_color[1:].upper())
    return FALLBACK_COLOR


def _tint(color: RGBColor, amount: float) -> RGBColor:
    """Mengt de kleur met wit; amount=0 is de kleur zelf, 1 is wit."""
    return RGBColor(*(int(round(c + (255 - c) * amount)) for c in color))


def _text_on(color: RGBColor) -> RGBColor:
    """Wit op een donkere kleur, donker op een lichte (bv. #8FAADC)."""
    r, g, b = color
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return DARK if luminance > 0.62 else WHITE


def _clamp_per_row(value: Any) -> int:
    try:
        n = int(value)
    except (TypeError, ValueError):
        return PER_ROW_DEFAULT
    return max(PER_ROW_MIN, min(PER_ROW_MAX, n))


def _rounded(slide, left, top, width, height, fill: RGBColor | None, *, line: RGBColor | None = None,
             line_width: float = 1.0, dashed: bool = False, radius: float = 0.16):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    try:
        shape.adjustments[0] = radius
    except (IndexError, ValueError):  # pragma: no cover -- vorm zonder afrondingsinstelling
        pass
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width)
        if dashed:
            shape.line.dash_style = MSO_LINE.DASH
    return shape


def _shape_text(shape, text: str, *, size: float, bold: bool = False, color: RGBColor = DARK,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=Inches(0.05)):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = margin
    tf.margin_right = margin
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = 'Calibri'


def _fit(text: str, width_emu: int, height_emu: int, size: float) -> str:
    """Kapt tekst af op wat er ongeveer in een vak van deze maat past
    (schatting op basis van de lettergrootte; liever iets te kort dan dat
    tekst buiten het vak loopt)."""
    per_line = max(4, int((width_emu / 914400) * 72 / (size * 0.52)))
    lines = max(1, int((height_emu / 914400) * 72 / (size * 1.22)))
    return _truncate(text, per_line * lines)


def _header(slide, kicker: str, title: str):
    title = _clean(title)
    size = 26 if len(title) <= 52 else 20
    _add_text(slide, MARGIN, Inches(0.45), CONTENT_W, Inches(0.3), _truncate(kicker.upper(), 110), size=12, bold=True, color=ACCENT)
    _add_text(slide, MARGIN, Inches(0.75), CONTENT_W, Inches(0.6), _truncate(title, 52 if size == 26 else 84), size=size, bold=True, color=DARK)
    _add_rect(slide, MARGIN, Inches(1.35), CONTENT_W, Emu(1), MUTED)


def _footer(slide, meta: dict[str, Any], page: int):
    parts = [_clean(meta.get('doelenboom')), _clean(meta.get('tenant'))]
    text = ' — '.join(p for p in parts if p)
    exported = _fmt_date(meta.get('exportedAt'))
    if exported != '-':
        text = f'{text} · {exported}' if text else exported
    _add_text(slide, MARGIN, SLIDE_H - Inches(0.4), CONTENT_W - Inches(0.6), Inches(0.3), _truncate(text, 150), size=9, color=MUTED)
    _add_text(
        slide, SLIDE_W - MARGIN - Inches(0.6), SLIDE_H - Inches(0.4), Inches(0.6), Inches(0.3),
        str(page), size=9, color=MUTED, align=PP_ALIGN.RIGHT,
    )


# ---- model: kolommen, zichtbare elementen en doorgetrokken relaties ------

class _Model:
    """De boom zoals de presentatie hem toont: alleen zichtbare kolommen en
    de elementen daarin, met relaties doorgetrokken over verborgen kolommen."""

    def __init__(self, data: dict[str, Any], options: dict[str, Any]):
        raw_columns = [c for c in (data.get('columns') or []) if isinstance(c, dict) and _clean(c.get('typeName'))]
        raw_columns.sort(key=lambda c: c.get('position') if isinstance(c.get('position'), int) else 0)
        all_names = [_clean(c.get('typeName')) for c in raw_columns]

        visible_opt = options.get('visibleColumns')
        if isinstance(visible_opt, list):
            wanted = {v for v in visible_opt if isinstance(v, str)}
            visible_names = [n for n in all_names if n in wanted]
        else:
            visible_names = list(all_names)
        if not visible_names:
            raise TreePptxError('Er is geen enkele zichtbare kolom om te exporteren.')

        slide_opt = options.get('slideColumns')
        slide_wanted = {v for v in slide_opt if isinstance(v, str)} if isinstance(slide_opt, list) else set()

        self.per_row = _clamp_per_row(options.get('perRow'))
        # Positie in de volledige kolomlijst, om te weten of twee zichtbare
        # kolommen ook echt buren zijn (alleen dan klopt het relatielabel).
        self.columns: list[dict[str, Any]] = []
        type_to_col: dict[str, int] = {}
        all_type_to_original: dict[str, int] = {}
        for original_index, c in enumerate(raw_columns):
            name = _clean(c.get('typeName'))
            aliases = [a for a in (c.get('aliases') or []) if isinstance(a, dict) and _clean(a.get('typeName'))]
            all_type_to_original[name] = original_index
            for a in aliases:
                all_type_to_original[_clean(a.get('typeName'))] = original_index
            if name not in visible_names:
                continue
            index = len(self.columns)
            type_to_col[name] = index
            alias_colors = {}
            for a in aliases:
                type_to_col[_clean(a.get('typeName'))] = index
                if a.get('color'):
                    alias_colors[_clean(a.get('typeName'))] = _rgb(a.get('color'))
            self.columns.append({
                'typeName': name,
                'title': _clean(c.get('title')) or name,
                'subtitle': _clean(c.get('subtitle')),
                'color': _rgb(c.get('color')),
                'aliasColors': alias_colors,
                'relationLabel': _clean(c.get('relationLabelToNext')),
                'originalIndex': original_index,
                'slides': name in slide_wanted,
                'elements': [],
            })

        # Elementen in boomvolgorde (zoals aangeleverd: sort_order, code).
        self.elements: dict[str, dict[str, Any]] = {}
        self.order: dict[str, int] = {}
        hidden_codes: set[str] = set()
        for raw in data.get('elements') or []:
            if not isinstance(raw, dict):
                continue
            code = _clean(raw.get('code'))
            if not code or code in self.elements or code in hidden_codes:
                continue
            el_type = _clean(raw.get('type'))
            col = type_to_col.get(el_type)
            if col is None:
                # Verborgen kolom (of onbekend type): niet tonen, maar wel
                # onthouden om relaties erdoorheen te kunnen doortrekken.
                hidden_codes.add(code)
                continue
            self.order[code] = len(self.order)
            self.elements[code] = {
                'code': code,
                'type': el_type,
                'col': col,
                'name': _clean(raw.get('name')),
                'description': _clean(raw.get('description')),
                'kpi': _clean(raw.get('kpi')),
                'taakveld': _clean(raw.get('taakveld')),
                'subtaakveld': _clean(raw.get('subtaakveld')),
            }
            self.columns[col]['elements'].append(code)

        self._visible_edges(data.get('edges') or [], hidden_codes)

        tag_names = {
            _clean(t.get('code')): _clean(t.get('name')) or _clean(t.get('code'))
            for t in (data.get('tags') or []) if isinstance(t, dict)
        }
        org_names = {
            _clean(o.get('code')): _clean(o.get('name')) or _clean(o.get('code'))
            for o in (data.get('orgUnits') or []) if isinstance(o, dict)
        }
        element_tags = data.get('elementTags') if isinstance(data.get('elementTags'), dict) else {}
        ob_org = data.get('obOrg') if isinstance(data.get('obOrg'), dict) else {}
        for code, el in self.elements.items():
            codes = element_tags.get(code) or []
            el['tags'] = [tag_names.get(_clean(tc), _clean(tc)) for tc in codes if _clean(tc)] if isinstance(codes, list) else []
            orgs = []
            rels = ob_org.get(code) or []
            for rel in rels if isinstance(rels, list) else []:
                if not isinstance(rel, dict) or not _clean(rel.get('org')):
                    continue
                name = org_names.get(_clean(rel.get('org')), _clean(rel.get('org')))
                relatietype = _clean(rel.get('relatietype'))
                orgs.append(f'{name} ({relatietype})' if relatietype else name)
            el['orgs'] = orgs

    def _visible_edges(self, raw_edges: list[Any], hidden_codes: set[str]):
        """Bouwt parents/children tussen zichtbare elementen. Een relatie
        naar een element in een verborgen kolom wordt doorgetrokken naar de
        eerstvolgende zichtbare elementen daarboven; zo'n doorgetrokken
        relatie is alleen 'primair' als elke schakel dat is."""
        up: dict[str, list[tuple[str, bool]]] = {}
        for e in raw_edges:
            if not isinstance(e, dict):
                continue
            source, target = _clean(e.get('source')), _clean(e.get('target'))
            if not source or not target or source == target:
                continue
            known = lambda c: c in self.elements or c in hidden_codes  # noqa: E731
            if not known(source) or not known(target):
                continue
            up.setdefault(source, []).append((target, _clean(e.get('weight')) != 'ondersteunend'))

        self.parents: dict[str, dict[str, bool]] = {c: {} for c in self.elements}
        self.children: dict[str, dict[str, bool]] = {c: {} for c in self.elements}
        for code in self.elements:
            # Iteratief (geen recursie): een kwaadaardig diepe of cyclische
            # keten mag de service niet laten crashen.
            best: dict[str, bool] = {}
            seen: dict[str, bool] = {}
            stack = [(t, primary) for t, primary in up.get(code, [])]
            while stack:
                node, primary = stack.pop()
                if node in self.elements:
                    if node != code:
                        best[node] = best.get(node, False) or primary
                    continue
                if seen.get(node) is True or (node in seen and not primary):
                    continue
                seen[node] = primary
                for t, p in up.get(node, []):
                    stack.append((t, primary and p))
            for parent, primary in best.items():
                self.parents[code][parent] = primary
                self.children[parent][code] = primary

    def _closure(self, start: str, graph: dict[str, dict[str, bool]]) -> dict[str, int]:
        """Alle elementen bereikbaar vanaf start, met hun afstand."""
        dist: dict[str, int] = {}
        frontier = [start]
        d = 0
        while frontier:
            d += 1
            nxt = []
            for node in frontier:
                for other in graph.get(node, {}):
                    if other != start and other not in dist:
                        dist[other] = d
                        nxt.append(other)
            frontier = nxt
        return dist

    def chain(self, code: str) -> dict[str, int]:
        """Het hele pad van één element: alles erboven en alles eronder,
        met de afstand tot het element (het element zelf = 0)."""
        result = {code: 0}
        for other, d in self._closure(code, self.parents).items():
            result[other] = d
        for other, d in self._closure(code, self.children).items():
            result.setdefault(other, d)
        return result

    def sorted_codes(self, codes) -> list[str]:
        return sorted(codes, key=lambda c: self.order.get(c, 0))

    def color_of(self, el: dict[str, Any]) -> RGBColor:
        col = self.columns[el['col']]
        return col['aliasColors'].get(el['type'], col['color'])

    def slide_columns(self) -> list[dict[str, Any]]:
        return [c for c in self.columns if c['slides']]

    def planned_slides(self) -> int:
        """Snoer + per gekozen kolom een tussenslide en een slide per element
        (zonder vervolgslides) -- dezelfde telling als in het dialoogvenster."""
        snoer_rows = math.ceil(len(self.columns) / self.per_row)
        total = max(1, math.ceil(snoer_rows / SNOER_MAX_ROWS))
        for c in self.slide_columns():
            total += 1 + len(c['elements'])
        return total


# ---- slide 1: de kolommen als snoer --------------------------------------

def _element_count_text(n: int) -> str:
    return '1 element' if n == 1 else f'{n} elementen'


def _slides_snoer(prs, model: _Model, meta: dict[str, Any], page: int) -> int:
    per_row = model.per_row
    columns = model.columns
    rows = [columns[i:i + per_row] for i in range(0, len(columns), per_row)]
    pages = [rows[i:i + SNOER_MAX_ROWS] for i in range(0, len(rows), SNOER_MAX_ROWS)]
    any_slides = any(c['slides'] for c in columns)

    gap_x = Inches(0.6)
    box_w = int((CONTENT_W - (per_row - 1) * gap_x) / per_row)
    area_top = Inches(2.15)
    area_bottom = SLIDE_H - Inches(0.6)

    number = 0
    row_offset = 0
    for page_index, page_rows in enumerate(pages):
        slide = _blank_slide(prs)
        title = _clean(meta.get('doelenboom')) or 'Doelenboom'
        _header(slide, 'Doelenboom' + (' — vervolg' if page_index else ''), title)
        intro = 'De kolommen van deze doelenboom, in volgorde.'
        if any_slides:
            intro += ' De gemarkeerde kolommen zijn hierna per element uitgewerkt.'
        _add_text(slide, MARGIN, Inches(1.47), CONTENT_W, Inches(0.3), intro, size=12, color=MUTED)

        n_rows = len(page_rows)
        gap_y = Inches(0.5)
        box_h = int(min(Inches(1.9), (area_bottom - area_top - (n_rows - 1) * gap_y) / n_rows))
        compact = box_h < Inches(1.25)
        band_h = Inches(0.36) if compact else Inches(0.46)
        title_size = 11 if (compact or box_w < Inches(2.2)) else 13

        for r, row in enumerate(page_rows):
            absolute_row = row_offset + r
            reverse = absolute_row % 2 == 1  # slingerend: oneven regels van rechts naar links
            top = int(area_top + r * (box_h + gap_y))
            for j, col in enumerate(row):
                number += 1
                slot = (per_row - 1 - j) if reverse else j
                left = int(MARGIN + slot * (box_w + gap_x))
                color = col['color']
                chosen = col['slides']

                body = _rounded(slide, left, top, box_w, box_h, _tint(color, 0.86),
                                line=DARK if chosen else color, line_width=2.25 if chosen else 1.0, radius=0.07)
                body.name = f'snoer-kolom-{number}'
                band = _add_rect(slide, left + Emu(12700), top + Emu(12700), box_w - Emu(25400), band_h, color)
                _shape_text(band, _fit(f'{number}. {col["title"]}', box_w - Inches(0.2), band_h, title_size),
                            size=title_size, bold=True, color=_text_on(color), align=PP_ALIGN.LEFT, margin=Inches(0.1))

                count_text = _element_count_text(len(col['elements']))
                if chosen:
                    count_text += ' · uitgewerkt' if box_w < Inches(2.2) else ' · per element uitgewerkt'
                sub_size = 9 if compact or box_w < Inches(2.2) else 10.5
                sub_top = top + band_h + Inches(0.07)
                if compact and box_w >= Inches(3.2):
                    # Lage, brede vakken: het aantal rechts in de gekleurde
                    # band, zodat de omschrijving de rest van het vak krijgt.
                    _add_text(slide, left + box_w - Inches(3.0), top, Inches(2.9), band_h, count_text, size=9,
                              bold=chosen, color=_text_on(color), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
                    sub_h = box_h - band_h - Inches(0.12)
                else:
                    count_h = Inches(0.24)
                    sub_h = box_h - band_h - count_h - Inches(0.14)
                    _add_text(slide, left + Inches(0.1), top + box_h - count_h - Inches(0.04), box_w - Inches(0.2), count_h,
                              count_text, size=9, bold=chosen,
                              color=DARK if chosen else MUTED, anchor=MSO_ANCHOR.BOTTOM)
                if col['subtitle'] and sub_h > Inches(0.15):
                    _add_text(slide, left + Inches(0.1), sub_top, box_w - Inches(0.2), sub_h,
                              _fit(col['subtitle'], box_w - Inches(0.2), sub_h, sub_size), size=sub_size, color=DARK)

                # Pijl naar de volgende kolom.
                global_index = number  # index van de volgende kolom in model.columns
                if global_index >= len(columns):
                    continue
                nxt = columns[global_index]
                label = col['relationLabel'] if nxt['originalIndex'] == col['originalIndex'] + 1 else ''
                if j < len(row) - 1:
                    arrow_w, arrow_h = Inches(0.36), Inches(0.24)
                    gap_left = left - gap_x if reverse else left + box_w
                    ax = int(gap_left + (gap_x - arrow_w) / 2)
                    ay = int(top + (box_h - arrow_h) / 2)
                    arrow = slide.shapes.add_shape(
                        MSO_SHAPE.LEFT_ARROW if reverse else MSO_SHAPE.RIGHT_ARROW, ax, ay, arrow_w, arrow_h)
                    _style_arrow(arrow)
                    if label:
                        label_w = Inches(1.7)
                        _add_text(slide, int(gap_left + gap_x / 2 - label_w / 2), top - Inches(0.27), label_w, Inches(0.24),
                                  _truncate(label, 30), size=9, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
                elif r < n_rows - 1:
                    # Laatste vak van de regel: omlaag naar de volgende regel.
                    arrow_w, arrow_h = Inches(0.24), Inches(0.34)
                    ax = int(left + (box_w - arrow_w) / 2)
                    ay = int(top + box_h + (gap_y - arrow_h) / 2)
                    arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, ax, ay, arrow_w, arrow_h)
                    _style_arrow(arrow)
                    if label:
                        # Label aan de buitenkant van de pijl (richting de
                        # slide-rand): daar staat nooit een ander label.
                        label_w = int(box_w / 2 - arrow_w / 2 - Inches(0.08) + Inches(0.5))
                        if reverse:
                            _add_text(slide, ax - Inches(0.08) - label_w, ay, label_w, arrow_h, _truncate(label, 30),
                                      size=9, color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
                        else:
                            _add_text(slide, ax + arrow_w + Inches(0.08), ay, label_w, arrow_h, _truncate(label, 30),
                                      size=9, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
        row_offset += n_rows
        _footer(slide, meta, page)
        page += 1
    return page


def _style_arrow(arrow):
    arrow.shadow.inherit = False
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = LINE_COLOR
    arrow.line.fill.background()


# ---- tussenslide per gekozen kolom ---------------------------------------

def _slide_section(prs, model: _Model, col: dict[str, Any], meta: dict[str, Any], page: int) -> int:
    slide = _blank_slide(prs)
    color = col['color']
    on = _text_on(color)
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, LIGHT_BG)
    block_top, block_h = Inches(2.1), Inches(2.9)
    _add_rect(slide, 0, block_top, SLIDE_W, block_h, color)
    _add_text(slide, MARGIN, block_top + Inches(0.35), CONTENT_W, Inches(0.35), 'KOLOM', size=13, bold=True, color=on)
    _add_text(slide, MARGIN, block_top + Inches(0.75), CONTENT_W, Inches(0.9), _truncate(col['title'], 60),
              size=40, bold=True, color=on)
    if col['subtitle']:
        _add_text(slide, MARGIN, block_top + Inches(1.7), CONTENT_W, Inches(0.8), _truncate(col['subtitle'], 220),
                  size=16, color=on)
    _add_text(slide, MARGIN, block_top + block_h + Inches(0.25), CONTENT_W, Inches(0.4),
              _element_count_text(len(col['elements'])), size=16, bold=True, color=DARK)
    _footer(slide, meta, page)
    return page + 1


# ---- element-slide: keten -------------------------------------------------

CHAIN_TOP = Inches(1.5)
CHAIN_HEAD_H = Inches(0.26)
CHAIN_NODES_TOP = Inches(1.84)
CHAIN_NODES_H = Inches(2.36)
TEXT_TOP = Inches(4.55)
TEXT_BOTTOM = SLIDE_H - Inches(0.55)


def _draw_chain(slide, model: _Model, code: str):
    chain = model.chain(code)
    n_cols = len(model.columns)
    gap = Inches(0.14) if n_cols <= 8 else Inches(0.08)
    col_w = int((CONTENT_W - (n_cols - 1) * gap) / n_cols)

    per_col: dict[int, list[str]] = {}
    for other in chain:
        per_col.setdefault(model.elements[other]['col'], []).append(other)

    # Per kolom: dichtst bij het element eerst houden, daarna boomvolgorde.
    shown: dict[int, list[str]] = {}
    more: dict[int, int] = {}
    for col, codes in per_col.items():
        if len(codes) > CHAIN_MAX_ROWS:
            keep = sorted(codes, key=lambda c: (chain[c], model.order.get(c, 0)))[:CHAIN_MAX_ROWS - 1]
            shown[col] = model.sorted_codes(keep)
            more[col] = len(codes) - len(keep)
        else:
            shown[col] = model.sorted_codes(codes)

    max_rows = max((len(shown[c]) + (1 if c in more else 0)) for c in shown)
    slot_h = int(min(Inches(0.66), CHAIN_NODES_H / max_rows))
    node_h = int(slot_h - Inches(0.07))
    size = 8 if slot_h >= Inches(0.5) else 7

    # Kolomkoppen.
    for index, col in enumerate(model.columns):
        left = int(MARGIN + index * (col_w + gap))
        head = _add_rect(slide, left, CHAIN_TOP, col_w, CHAIN_HEAD_H, col['color'] if index in shown else _tint(col['color'], 0.7))
        _shape_text(head, _fit(col['title'], col_w, CHAIN_HEAD_H, 8), size=8, bold=True,
                    color=_text_on(col['color']) if index in shown else DARK, margin=Inches(0.03))

    # Posities van de vakken (verticaal gecentreerd per kolom).
    pos: dict[str, tuple[int, int]] = {}
    more_pos: dict[int, tuple[int, int]] = {}
    for col, codes in shown.items():
        count = len(codes) + (1 if col in more else 0)
        start = int(CHAIN_NODES_TOP + (CHAIN_NODES_H - count * slot_h) / 2)
        left = int(MARGIN + col * (col_w + gap))
        for i, other in enumerate(codes):
            pos[other] = (left, start + i * slot_h)
        if col in more:
            more_pos[col] = (left, start + len(codes) * slot_h)

    # Eerst de lijnen, zodat de vakken er bovenop liggen.
    has_supporting = False
    for child in pos:
        for parent, primary in model.parents[child].items():
            if parent not in pos:
                continue
            (cl, ct), (pl, pt) = pos[child], pos[parent]
            if cl == pl:
                continue  # zelfde kolom: geen zinvolle lijn te trekken
            if cl < pl:
                x1, x2 = cl + col_w, pl
            else:
                x1, x2 = cl, pl + col_w
            y1, y2 = ct + node_h // 2, pt + node_h // 2
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
            line.line.color.rgb = LINE_COLOR
            line.line.width = Pt(1.25 if primary else 1.0)
            line.name = f'relatie {child} > {parent}' + ('' if primary else ' (ondersteunend)')
            if not primary:
                line.line.dash_style = MSO_LINE.DASH
                has_supporting = True

    for other, (left, top) in pos.items():
        el = model.elements[other]
        color = model.color_of(el)
        label = _fit(f'{other} {el["name"]}'.strip(), col_w - Inches(0.08), node_h, size)
        if other == code:
            node = _rounded(slide, left, top, col_w, node_h, color, line=DARK, line_width=2.25)
            _shape_text(node, label, size=size, bold=True, color=_text_on(color), margin=Inches(0.04))
            node.name = f'element {other} (geselecteerd)'
        else:
            node = _rounded(slide, left, top, col_w, node_h, _tint(color, 0.8), line=color, line_width=0.75)
            _shape_text(node, label, size=size, color=DARK, margin=Inches(0.04))
            node.name = f'element {other}'
    for col, (left, top) in more_pos.items():
        node = _rounded(slide, left, top, col_w, node_h, None, line=LINE_COLOR, line_width=0.75, dashed=True)
        _shape_text(node, f'+ {more[col]} meer', size=size, color=MUTED, margin=Inches(0.04))

    legend = 'Het gekleurde vak met donkere rand is dit element.'
    if len(pos) > 1:
        legend += ' Doorgetrokken lijn = primaire relatie'
        legend += ', stippellijn = ondersteunende relatie.' if has_supporting else '.'
    else:
        legend = 'Dit element heeft (nog) geen verbindingen met andere zichtbare elementen.'
    _add_text(slide, MARGIN, CHAIN_NODES_TOP + CHAIN_NODES_H + Inches(0.04), CONTENT_W, Inches(0.2), legend, size=8, color=MUTED)


# ---- element-slide: tekst --------------------------------------------------

LEFT_W = Inches(6.7)
RIGHT_LEFT = MARGIN + LEFT_W + Inches(0.5)
RIGHT_W = CONTENT_W - LEFT_W - Inches(0.5)
LEFT_CHARS = 88
RIGHT_CHARS = 60
BODY_SIZE = 11
LINE_H = Inches(0.2)
HEADING_COST = 1.5  # kopje + witruimte ervoor, in regels


def _lines_for(text: str, chars: int) -> list[str]:
    out: list[str] = []
    for para in text.split('\n'):
        para = para.strip()
        if not para:
            continue
        out.extend(textwrap.wrap(para, chars) or [''])
    return out


def _paginate_blocks(blocks: list[tuple[str, list[str]]], chars: int, capacities: list[float]) -> list[list[tuple[str, str]]]:
    """Verdeelt blokken (kopje + alinea's) over pagina's met elk een
    capaciteit in regels. Geeft per pagina een lijst ('h'|'p', tekst).
    Wat op de laatste pagina niet meer past wordt ingekort met '…'."""
    pages: list[list[tuple[str, str]]] = [[]]
    used = 0.0

    def capacity() -> float:
        return capacities[min(len(pages) - 1, len(capacities) - 1)]

    def last_page() -> bool:
        return len(pages) >= len(capacities)

    truncated = False
    for heading, paragraphs in blocks:
        if truncated:
            break
        first = True
        for para in paragraphs:
            lines = _lines_for(para, chars)
            while lines and not truncated:
                need_heading = HEADING_COST if first else 0
                room = int(capacity() - used - need_heading)
                if room < 1 or (first and room < min(2, len(lines))):
                    if last_page():
                        truncated = True
                        break
                    pages.append([])
                    used = 0.0
                    if not first:
                        pages[-1].append(('h', f'{heading} (vervolg)'))
                        used += HEADING_COST
                    continue
                if first:
                    pages[-1].append(('h', heading))
                    used += HEADING_COST
                    first = False
                take = lines[:room]
                lines = lines[room:]
                text = ' '.join(take)
                if lines and last_page():
                    text = text.rstrip() + ' …'
                    lines = []
                    truncated = True
                pages[-1].append(('p', text))
                used += len(take)
                if lines:
                    pages.append([])
                    used = 0.0
                    pages[-1].append(('h', f'{heading} (vervolg)'))
                    used += HEADING_COST
    if truncated and pages[-1]:
        kind, text = pages[-1][-1]
        if kind == 'p' and not text.endswith('…'):
            pages[-1][-1] = (kind, text + ' …')
    return [p for p in pages if p] or [[]]


def _draw_text_column(slide, left, top, width, height, items: list[tuple[str, str]]):
    if not items:
        return
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for kind, text in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        run = p.add_run()
        run.font.name = 'Calibri'
        if kind == 'h':
            run.text = text.upper()
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = ACCENT
            if not first:
                p.space_before = Pt(7)
        else:
            run.text = text
            run.font.size = Pt(BODY_SIZE)
            run.font.color.rgb = DARK
        first = False


def _relation_lines(model: _Model, related: dict[str, bool]) -> list[str]:
    lines = []
    for other in model.sorted_codes(related):
        el = model.elements[other]
        line = f'• {other} — {el["name"]}' if el['name'] else f'• {other}'
        line = _truncate(line, RIGHT_CHARS * 2 - 20)
        line += f' ({model.columns[el["col"]]["title"]}'
        line += ')' if related[other] else ', ondersteunend)'
        lines.append(line)
    return lines


def _element_blocks(model: _Model, el: dict[str, Any], title_truncated: bool):
    left: list[tuple[str, list[str]]] = []
    description = el['description']
    if title_truncated and el['name']:
        description = f'{el["name"]}\n{description}' if description else el['name']
    # Elke regel uit de bron is een eigen alinea (anders lopen ze in elkaar over).
    paragraphs = [p.strip() for p in description.split('\n') if p.strip()]
    left.append(('Beschrijving', paragraphs or ['Geen beschrijving ingevuld.']))
    if el['kpi']:
        left.append(('KPI / indicator', [p.strip() for p in el['kpi'].split('\n') if p.strip()]))

    right: list[tuple[str, list[str]]] = []
    parents = model.parents[el['code']]
    children = model.children[el['code']]
    if parents:
        right.append((f'Verbonden met — bovenliggend ({len(parents)})', _relation_lines(model, parents)))
    if children:
        right.append((f'Verbonden met — onderliggend ({len(children)})', _relation_lines(model, children)))
    if not parents and not children:
        right.append(('Verbonden met', ['Geen verbindingen.']))
    if el['tags']:
        right.append(('Tags', [', '.join(el['tags'])]))
    if el['orgs']:
        right.append(('Organisatieonderdelen', ['; '.join(el['orgs'])]))
    return left, right


def _slides_element(prs, model: _Model, el: dict[str, Any], meta: dict[str, Any], page: int) -> int:
    col = model.columns[el['col']]
    kicker = col['title']
    if el['type'] != col['typeName']:
        kicker += f' ({el["type"]})'
    veld = ' / '.join(v for v in (el['taakveld'], el['subtaakveld']) if v)
    if veld:
        kicker += f' · {veld}'
    title = f'{el["code"]} — {el["name"]}' if el['name'] else el['code']
    title_truncated = len(title) > 84

    first_cap = (TEXT_BOTTOM - TEXT_TOP) / LINE_H
    next_cap = (TEXT_BOTTOM - Inches(1.6)) / LINE_H
    capacities = [first_cap] + [next_cap] * (MAX_SLIDES_PER_ELEMENT - 1)
    left_blocks, right_blocks = _element_blocks(model, el, title_truncated)
    left_pages = _paginate_blocks(left_blocks, LEFT_CHARS, capacities)
    right_pages = _paginate_blocks(right_blocks, RIGHT_CHARS, capacities)

    for index in range(max(len(left_pages), len(right_pages))):
        slide = _blank_slide(prs)
        _header(slide, kicker + (' — vervolg' if index else ''), title)
        if index == 0:
            _draw_chain(slide, model, el['code'])
            top = TEXT_TOP
        else:
            top = Inches(1.6)
        height = TEXT_BOTTOM - top
        if index < len(left_pages):
            _draw_text_column(slide, MARGIN, top, LEFT_W, height, left_pages[index])
        if index < len(right_pages):
            _draw_text_column(slide, RIGHT_LEFT, top, RIGHT_W, height, right_pages[index])
        _footer(slide, meta, page)
        page += 1
    return page


# ---- ingang ----------------------------------------------------------------

def build_tree_pptx(data: dict[str, Any], options: dict[str, Any] | None, meta: dict[str, Any] | None) -> bytes:
    """Bouwt de doelenboom-presentatie. data = de boom (columns, elements,
    edges, tags, elementTags, orgUnits, obOrg), options = {visibleColumns,
    slideColumns, perRow}, meta = {doelenboom, tenant, exportedAt}.
    Gooit TreePptxError bij een ongeldige keuze."""
    data = data if isinstance(data, dict) else {}
    options = options if isinstance(options, dict) else {}
    meta = meta if isinstance(meta, dict) else {}

    model = _Model(data, options)
    planned = model.planned_slides()
    if planned > MAX_SLIDES:
        raise TreePptxError(
            f'Deze keuze levert {planned} slides op; het maximum is {MAX_SLIDES}. '
            'Kies minder kolommen om per element uit te werken.'
        )

    prs = _new_presentation()
    page = _slides_snoer(prs, model, meta, 1)
    for col in model.slide_columns():
        page = _slide_section(prs, model, col, meta, page)
        for code in col['elements']:
            page = _slides_element(prs, model, model.elements[code], meta, page)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
