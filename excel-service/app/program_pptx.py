"""
Programma-PowerPoint (DOEL-112): één presentatie voor MEERDERE projecten
tegelijk, bedoeld voor een programmamanager / stuurgroep. Hergebruikt de
huisstijl en de voortgangsslide van de project-PowerPoint (project_pptx.py).

Opbouw:
  1.  Tijdlijnen -- de gekozen projecten MET datums op één gezamenlijke
      tijdas, per project de looptijd, de fasen (samenvattende activiteiten)
      en de mijlpalen, met een peildatum-lijn en de RAG-status als kleur.
      Altijd precies één slide: de rijen schalen mee met het aantal
      projecten (snel overzicht; de detailplanning staat op de
      voortgangsslide per project). Dit is de "Home"-slide.
  1b. Alle projecten -- ALLE gekozen projecten (ook die nog zonder planning)
      als compact raster op één slide: RAG-kleur, naam, code en de stand
      (open, te laat, eerstvolgende datum) of "Nog geen planning".
  2.  Deliverables -- alle nog niet opgeleverde deliverables/mijlpalen van de
      gekozen projecten, in vier blokken: "Te laat" (verwachte datum vóór de
      peildatum), "Komende N maanden" (op datum), "Later" (op datum) en
      "Zonder datum". Loopt door op vervolgslides.
  3.  Bijdrage aan doelen -- per project aan welke elementen in de doelenboom
      het direct bijdraagt, en welke doelen het uiteindelijk dient (de
      bovenste kolommen worden door de API berekend, zie
      api/src/routes/programPptx.ts).
  4.  Voortgang per project -- de overzichtsslide uit de project-PowerPoint
      (status/RAG, projecttijdlijn, aandachtspunten), projecten alfabetisch.

Navigatie: een projectnaam of -balk op de tijdlijn, en een project of
deliverable op de deliverables-/doelenslides springt naar de voortgangsslide
van dat project. Elke slide behalve de eerste heeft een Home-knop terug naar
de (eerste) tijdlijnslide -- zelfde opzet als DOEL-100/DOEL-106.

Aangeroepen door api/src/routes/programPptx.ts via POST /program-pptx.
"""
from __future__ import annotations

import calendar
import io
import re
from datetime import date
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .project_pptx import (
    ACCENT,
    CONTENT_W,
    DARK,
    LIGHT_BG,
    MAANDEN_KORT,
    MARGIN,
    MUTED,
    RAG_COLORS,
    SLIDE_H,
    SLIDE_W,
    TIMELINE_AXIS_COLOR,
    TIMELINE_DEADLINE_COLOR,
    WHITE,
    _add_rect,
    _add_text,
    _axis_bounds,
    _blank_slide,
    _fmt_date,
    _new_presentation,
    _parse_iso_date,
    _pct,
    _rag_color,
    _slide_header,
    _slide_overview,
    _truncate,
)

# Grenzen -- dezelfde waarden staan in api/src/routes/programPptx.ts (die
# valideert als eerste); hier nogmaals als verdedigingslinie.
MAX_PROJECTS = 50
MAX_SLIDES = 300
WINDOW_MIN = 1
WINDOW_MAX = 12
WINDOW_DEFAULT = 3

# Tijdlijn: hooguit deze rijhoogte; bij meer projecten krimpen de rijen zodat
# alles op één slide past (MAX_PROJECTS = 50 past nog leesbaar).
TIMELINE_MAX_ROW_H = Inches(0.53)
DELIVERABLE_ROWS_PER_SLIDE = 14
GOAL_ROWS_PER_SLIDE = 7

LATE_COLOR = RAG_COLORS['rood']
LATE_BG = RGBColor(0xFB, 0xE4, 0xE6)
SECTION_BG = RGBColor(0xE3, 0xEA, 0xF4)
PHASE_COLORS = [RGBColor(0x2F, 0x55, 0x97), RGBColor(0x5B, 0x7C, 0xB8)]
SPAN_COLOR = RGBColor(0xB7, 0xC6, 0xE6)
MILESTONE_COLOR = RGBColor(0x00, 0x28, 0x55)

_CONTROL_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]')
_ISO_RE = re.compile(r'^\d{4}-\d{2}-\d{2}$')


class ProgramPptxError(ValueError):
    """Ongeldige invoer of keuze -- main.py maakt hier een 422 van."""


# ---- invoer opschonen en valideren --------------------------------------

def _clean_data(value: Any, depth: int = 0) -> Any:
    """Haalt besturingstekens (die niet in XML mogen) uit alle strings in de
    invoer, recursief. Diepte begrensd: de invoer komt van de eigen API,
    maar wordt hier niet blind vertrouwd."""
    if depth > 8:
        return None
    if isinstance(value, str):
        return _CONTROL_RE.sub('', value)
    if isinstance(value, list):
        return [_clean_data(v, depth + 1) for v in value]
    if isinstance(value, dict):
        return {str(k): _clean_data(v, depth + 1) for k, v in value.items()}
    return value


def add_months_same_day(d: date, months: int) -> date:
    """13 okt + 3 mnd = 13 jan; 31 jan + 1 mnd = 28/29 feb (geklemd)."""
    total = d.year * 12 + (d.month - 1) + months
    y, m = divmod(total, 12)
    m += 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def parse_options(options: dict[str, Any] | None, meta: dict[str, Any] | None) -> tuple[date, int]:
    """(peildatum, venster in maanden). Ontbrekend = vandaag (meta.exportedAt)
    en 3 maanden; ongeldig = fout (de API heeft dit al afgevangen)."""
    options = options if isinstance(options, dict) else {}
    meta = meta if isinstance(meta, dict) else {}
    ref_raw = options.get('referenceDate')
    if ref_raw in (None, ''):
        ref = _parse_iso_date(meta.get('exportedAt')) or date.today()
    else:
        if not isinstance(ref_raw, str) or not _ISO_RE.match(ref_raw):
            raise ProgramPptxError('De peildatum moet de vorm JJJJ-MM-DD hebben.')
        ref = _parse_iso_date(ref_raw)
        if ref is None or not (2000 <= ref.year <= 2100):
            raise ProgramPptxError('De peildatum is geen geldige datum.')
    window = options.get('windowMonths', WINDOW_DEFAULT)
    if isinstance(window, bool) or not isinstance(window, int) or not (WINDOW_MIN <= window <= WINDOW_MAX):
        raise ProgramPptxError(f'Het venster moet tussen {WINDOW_MIN} en {WINDOW_MAX} maanden liggen.')
    return ref, window


def _projects_from(data: Any) -> list[dict[str, Any]]:
    if not isinstance(data, dict):
        raise ProgramPptxError('Geen projectgegevens ontvangen.')
    raw = data.get('projects')
    if not isinstance(raw, list) or not raw:
        raise ProgramPptxError('Kies minstens één project.')
    if len(raw) > MAX_PROJECTS:
        raise ProgramPptxError(f'Kies hooguit {MAX_PROJECTS} projecten.')
    projects = []
    for entry in raw:
        if not isinstance(entry, dict) or not isinstance(entry.get('project'), dict):
            raise ProgramPptxError('Ongeldige projectgegevens ontvangen.')
        p = dict(entry)
        for key in ('products', 'activities', 'productDependencies', 'activityDependencies'):
            p[key] = [x for x in (p.get(key) or []) if isinstance(x, dict)] if isinstance(p.get(key), list) else []
        p['goals'] = p.get('goals') if isinstance(p.get('goals'), dict) else {}
        projects.append(p)
    return projects


def sort_projects(projects: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Alfabetisch op naam (hoofdletterongevoelig), dan op code."""
    def key(p):
        proj = p['project']
        return (str(proj.get('name') or '').casefold(), str(proj.get('code') or ''))
    return sorted(projects, key=key)


# ---- deliverables indelen -------------------------------------------------

SECTION_KEYS = ('late', 'window', 'later', 'undated')


def classify_deliverables(projects: list[dict[str, Any]], ref: date, window_months: int) -> dict[str, list[dict[str, Any]]]:
    """Alle NIET opgeleverde producten (geen werkelijke datum) van de gekozen
    projecten, ingedeeld op hun verwachte datum t.o.v. de peildatum:
      late    -- verwachte datum vóór de peildatum
      window  -- peildatum t/m (peildatum + N maanden), exclusief het eind
      later   -- daarna
      undated -- geen (geldige) verwachte datum
    Binnen een blok op datum, dan projectnaam, dan deliverable-naam."""
    window_end = add_months_same_day(ref, window_months)
    sections: dict[str, list[dict[str, Any]]] = {k: [] for k in SECTION_KEYS}
    for p in projects:
        proj = p['project']
        for prod in p['products']:
            if prod.get('werkelijkeDatum'):
                continue
            d = _parse_iso_date(prod.get('verwachteDatum'))
            row = {
                'projectCode': str(proj.get('code') or ''),
                'projectName': str(proj.get('name') or ''),
                'name': str(prod.get('name') or ''),
                'type': 'Mijlpaal' if prod.get('type') == 'mijlpaal' else 'Deliverable',
                'date': d,
                'deadline': _parse_iso_date(prod.get('deadline')),
                'pct': _pct(prod.get('pctGereed')),
            }
            if d is None:
                sections['undated'].append(row)
            elif d < ref:
                sections['late'].append(row)
            elif d < window_end:
                sections['window'].append(row)
            else:
                sections['later'].append(row)
    by_date = lambda r: (r['date'] or date.max, r['projectName'].casefold(), r['name'].casefold())  # noqa: E731
    for k in ('late', 'window', 'later'):
        sections[k].sort(key=by_date)
    sections['undated'].sort(key=lambda r: (r['projectName'].casefold(), r['name'].casefold()))
    return sections


def _short_range(start: date, end_excl: date) -> str:
    last = date.fromordinal(end_excl.toordinal() - 1)
    def fmt(d: date, with_year: bool) -> str:
        return f'{d.day} {MAANDEN_KORT[d.month - 1]}' + (f' {d.year}' if with_year else '')
    return f'{fmt(start, start.year != last.year)} – {fmt(last, True)}'


def section_titles(ref: date, window_months: int) -> dict[str, str]:
    end = add_months_same_day(ref, window_months)
    maanden = 'maand' if window_months == 1 else 'maanden'
    return {
        'late': f'Te laat — verwachte datum vóór {_fmt_date(ref.isoformat())}',
        'window': f'Komende {window_months} {maanden} — {_short_range(ref, end)}',
        'later': f'Later — vanaf {_fmt_date(end.isoformat())}',
        'undated': 'Nog geen datum',
    }


def paginate_deliverables(sections: dict[str, list[dict[str, Any]]], rows_per_slide: int = DELIVERABLE_ROWS_PER_SLIDE) -> list[list[tuple]]:
    """Verdeelt de blokken over slides. Elke slide-entry is
    ('section', key, continued) of ('row', key, row). Een bloktitel komt
    nooit als laatste regel van een slide (dan begint het blok op de
    volgende), en een blok dat doorloopt krijgt bovenaan de volgende slide
    zijn titel opnieuw, met "(vervolg)"."""
    pages: list[list[tuple]] = []
    page: list[tuple] = []
    for key in SECTION_KEYS:
        rows = sections.get(key) or []
        if not rows:
            continue
        if len(page) > rows_per_slide - 2:
            pages.append(page)
            page = []
        page.append(('section', key, False))
        for row in rows:
            if len(page) >= rows_per_slide:
                pages.append(page)
                page = [('section', key, True)]
            page.append(('row', key, row))
    if page:
        pages.append(page)
    return pages


# ---- links tussen slides --------------------------------------------------

HOME_KEY = ('home',)


class _Links:
    """Klikbare vormen; koppeling aan de doelslide gebeurt pas aan het eind,
    omdat de voortgangsslides achteraan in de presentatie komen."""

    def __init__(self):
        self.pending: list[tuple[Any, Any, str]] = []
        self.slides: dict[Any, Any] = {}

    def add(self, shape, key: Any, tooltip: str):
        self.pending.append((shape, key, _truncate(tooltip, 200)))

    def resolve(self):
        for shape, key, tooltip in self.pending:
            target = self.slides.get(key)
            if target is None:
                continue
            shape.click_action.target_slide = target
            hlink = shape._element.find('.//' + qn('a:hlinkClick'))
            if hlink is not None and tooltip:
                hlink.set('tooltip', tooltip)


def _project_key(code: str) -> tuple[str, str]:
    return ('project', code)


HOME_W = Inches(0.62)
HOME_H = Inches(0.24)
PAGE_NO_W = Inches(0.6)


def _home_button(slide, links: _Links):
    home = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, SLIDE_W - MARGIN - PAGE_NO_W - HOME_W, SLIDE_H - Inches(0.44), HOME_W, HOME_H,
    )
    home.shadow.inherit = False
    home.fill.solid()
    home.fill.fore_color.rgb = ACCENT
    home.line.fill.background()
    home.name = 'home-knop'
    tf = home.text_frame
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = 'Home'
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = 'Calibri'
    links.add(home, HOME_KEY, 'Terug naar de tijdlijnen')


def _footer(slide, meta: dict[str, Any], page: int, links: _Links | None):
    parts = ['Programma-overzicht', str(meta.get('doelenboom') or ''), str(meta.get('tenant') or '')]
    text = ' — '.join(p for p in parts if p)
    exported = _fmt_date(meta.get('exportedAt'))
    if exported != '-':
        text += f' · {exported}'
    text_w = CONTENT_W - PAGE_NO_W - (HOME_W + Inches(0.2) if links is not None else 0)
    _add_text(slide, MARGIN, SLIDE_H - Inches(0.4), text_w, Inches(0.3), _truncate(text, 150), size=9, color=MUTED)
    if links is not None:
        _home_button(slide, links)
    _add_text(
        slide, SLIDE_W - MARGIN - PAGE_NO_W, SLIDE_H - Inches(0.4), PAGE_NO_W, Inches(0.3),
        str(page), size=9, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def _page_suffix(i: int, n: int) -> str:
    return f' — pagina {i + 1}/{n}' if n > 1 else ''


# ---- slide 1: tijdlijnen --------------------------------------------------

def _phases(activities: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Fasen = samenvattende activiteiten (MS Project 'Summary', of handmatig
    zo gemarkeerd) met een start- én einddatum. Bij geneste fasen alleen het
    bovenste WBS-niveau, anders stapelen ze over elkaar."""
    summaries = []
    for a in activities:
        if not a.get('isSummary'):
            continue
        start = _parse_iso_date(a.get('startDate'))
        end = _parse_iso_date(a.get('endDate'))
        if start is None or end is None:
            continue
        if end < start:
            start, end = end, start
        wbs = str(a.get('wbs') or '')
        summaries.append({'name': str(a.get('name') or ''), 'start': start, 'end': end,
                          'depth': wbs.count('.') if wbs else None})
    depths = [s['depth'] for s in summaries if s['depth'] is not None]
    if depths:
        top = min(depths)
        summaries = [s for s in summaries if s['depth'] is None or s['depth'] == top]
    summaries.sort(key=lambda s: (s['start'], s['end']))
    return summaries


def _milestones(p: dict[str, Any], ref: date) -> list[dict[str, Any]]:
    """Mijlpalen: producten van type 'mijlpaal' (gehaald = werkelijke datum)
    en activiteiten die als mijlpaal zijn gemarkeerd (gehaald = vóór of op
    de peildatum)."""
    result = []
    for prod in p['products']:
        if prod.get('type') != 'mijlpaal':
            continue
        done = _parse_iso_date(prod.get('werkelijkeDatum'))
        d = done or _parse_iso_date(prod.get('verwachteDatum'))
        if d:
            result.append({'t': d, 'done': done is not None, 'name': str(prod.get('name') or '')})
    for a in p['activities']:
        if not a.get('isMilestone'):
            continue
        d = _parse_iso_date(a.get('startDate')) or _parse_iso_date(a.get('endDate'))
        if d:
            result.append({'t': d, 'done': d <= ref, 'name': str(a.get('name') or '')})
    result.sort(key=lambda m: m['t'])
    return result


def project_dates(p: dict[str, Any]) -> list[date]:
    dates: list[date] = []
    for a in p['activities']:
        for k in ('startDate', 'endDate'):
            d = _parse_iso_date(a.get(k))
            if d:
                dates.append(d)
    for prod in p['products']:
        for k in ('verwachteDatum', 'werkelijkeDatum', 'deadline'):
            d = _parse_iso_date(prod.get(k))
            if d:
                dates.append(d)
    return dates


LABEL_W = Inches(3.0)


def _fit_chars(width_emu: int, size_pt: float) -> int:
    """Grove schatting van het aantal tekens dat op één regel past."""
    return max(4, int((width_emu / 12700) / (size_pt * 0.5)))


def _slide_timeline(prs, projects_with_dates, all_dates: list[date], ref: date, subtitle: str,
                    links: _Links, ref_label: str):
    """Alle projecten met datums op één slide; de rijhoogte, letters, balken
    en mijlpalen schalen mee met het aantal projecten."""
    slide = _blank_slide(prs)
    _slide_header(slide, 'Programma', 'Tijdlijnen')
    _add_text(slide, MARGIN, Inches(1.48), CONTENT_W, Inches(0.3), subtitle, size=12, color=MUTED)

    if not projects_with_dates:
        _add_text(slide, MARGIN, Inches(2.2), CONTENT_W, Inches(0.5),
                  'Geen van de gekozen projecten heeft al een planning met datums. '
                  'Zie het overzicht van alle projecten op de volgende slide.', size=14, color=MUTED)
        return slide

    quarterly, bounds = _axis_bounds(all_dates, ref)
    axis_start, axis_end = bounds[0], bounds[-1]
    span_days = (axis_end - axis_start).days or 1
    inner_left = int(MARGIN + LABEL_W + Inches(0.1))
    inner_w = int(CONTENT_W - LABEL_W - Inches(0.2))

    def x_for(d: date) -> int:
        frac = min(max((d - axis_start).days / span_days, 0.0), 1.0)
        return int(inner_left + inner_w * frac)

    n = len(projects_with_dates)
    axis_top = int(Inches(2.1))
    rows_top = axis_top + int(Inches(0.06))
    legend_top = int(SLIDE_H - Inches(0.8))
    area = legend_top - int(Inches(0.08)) - rows_top
    row_h = min(int(TIMELINE_MAX_ROW_H), area // n)
    rows_bottom = rows_top + row_h * n
    scale = row_h / int(TIMELINE_MAX_ROW_H)
    row_pt = row_h / 12700
    label_size = max(6.0, min(11.0, round(row_pt * 0.55, 1)))
    two_lines = scale >= 0.8

    # Maandkoppen/kwartaalkoppen en verticale rasterlijnen.
    for i in range(len(bounds) - 1):
        gx = x_for(bounds[i])
        nx = x_for(bounds[i + 1])
        if i > 0:
            _add_rect(slide, gx, axis_top, Pt(0.75), rows_bottom - axis_top, TIMELINE_AXIS_COLOR)
        if nx - gx >= int(Inches(0.55)):
            label = (f'K{(bounds[i].month - 1) // 3 + 1} {bounds[i].year}' if quarterly
                     else f'{MAANDEN_KORT[bounds[i].month - 1].capitalize()} {bounds[i].year}')
            _add_text(slide, gx, axis_top - Inches(0.25), nx - gx, Inches(0.22), label, size=9, color=MUTED, align=PP_ALIGN.CENTER)
    _add_rect(slide, inner_left, axis_top, inner_w, Pt(1), TIMELINE_AXIS_COLOR)

    bar_h = max(int(Pt(3)), min(int(Inches(0.24)), int(row_h * 0.45)))
    thin_h = max(int(Pt(2)), bar_h // 3)
    ms_size = max(int(Pt(5)), min(int(Inches(0.17)), int(row_h * 0.6)))
    inset = min(int(Inches(0.08)), int(row_h * 0.15))
    label_w = int(LABEL_W - Inches(0.25))

    for i, p in enumerate(projects_with_dates):
        proj = p['project']
        code = str(proj.get('code') or '')
        name = str(proj.get('name') or code)
        top = rows_top + i * row_h
        mid = top + row_h // 2
        if i % 2 == 1:
            _add_rect(slide, MARGIN, top, CONTENT_W, row_h, LIGHT_BG)

        rag = ((proj.get('status') or {}) if isinstance(proj.get('status'), dict) else {}).get('rag')
        _add_rect(slide, MARGIN, top + inset, Inches(0.08), max(row_h - 2 * inset, int(Pt(2))), _rag_color(rag))
        max_chars = _fit_chars(label_w, label_size) * (2 if two_lines else 1) - (len(code) + 2 if code else 0)
        label = _add_text(
            slide, MARGIN + Inches(0.18), top, label_w, row_h,
            _truncate(name, max(6, max_chars)), size=label_size, bold=True, color=DARK,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        if code:
            run = label.text_frame.paragraphs[0].add_run()
            run.text = f'  {code}'
            run.font.size = Pt(max(5.5, label_size - 3))
            run.font.color.rgb = MUTED
            run.font.name = 'Calibri'
        links.add(label, _project_key(code), f'Naar de voortgang van {name}')

        dates = project_dates(p)
        phases = _phases(p['activities'])
        span_x0, span_x1 = x_for(min(dates)), x_for(max(dates))
        span_hh = thin_h if phases else bar_h
        span = _add_rect(slide, span_x0, mid - span_hh // 2, max(span_x1 - span_x0, int(Pt(3))), span_hh, SPAN_COLOR)
        links.add(span, _project_key(code), f'Naar de voortgang van {name}')

        for j, ph in enumerate(phases):
            x0, x1 = x_for(ph['start']), x_for(ph['end'])
            w = max(x1 - x0, int(Pt(3)))
            bar = _add_rect(slide, x0, mid - bar_h // 2, w, bar_h, PHASE_COLORS[j % len(PHASE_COLORS)])
            phase_size = min(8.0, label_size - 1)
            if w >= int(Inches(0.6)) and bar_h >= int(Pt(phase_size + 2)):
                tf = bar.text_frame
                tf.margin_left = tf.margin_right = Inches(0.04)
                tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.word_wrap = False
                para = tf.paragraphs[0]
                para.alignment = PP_ALIGN.LEFT
                run = para.add_run()
                run.text = _truncate(ph['name'], _fit_chars(w - int(Inches(0.08)), phase_size))
                run.font.size = Pt(phase_size)
                run.font.color.rgb = WHITE
                run.font.name = 'Calibri'
            links.add(bar, _project_key(code), f'{ph["name"]} ({_fmt_date(ph["start"].isoformat())} – {_fmt_date(ph["end"].isoformat())})')

        # Ruim genoeg: mijlpaal boven de balk; compact: op de balk zelf.
        lift = int(Inches(0.2) * scale) if scale >= 0.75 else 0
        for m in _milestones(p, ref):
            cx = x_for(m['t'])
            shape = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, cx - ms_size // 2, mid - ms_size // 2 - lift, ms_size, ms_size)
            shape.shadow.inherit = False
            shape.fill.solid()
            shape.fill.fore_color.rgb = MILESTONE_COLOR if m['done'] else WHITE
            shape.line.color.rgb = MILESTONE_COLOR
            shape.line.width = Pt(1.25 if scale >= 0.5 else 0.75)
            links.add(shape, _project_key(code), f'Mijlpaal: {m["name"]} ({_fmt_date(m["t"].isoformat())})')

    if axis_start <= ref <= axis_end:
        rx = x_for(ref)
        _add_rect(slide, rx, axis_top - Inches(0.05), Pt(1.5), rows_bottom - axis_top + Inches(0.05), TIMELINE_DEADLINE_COLOR)
        _add_text(slide, rx - Inches(0.5), axis_top - Inches(0.47), Inches(1.0), Inches(0.2), ref_label,
                  size=8, bold=True, color=TIMELINE_DEADLINE_COLOR, align=PP_ALIGN.CENTER)

    _rag_legend(slide, legend_top, ref_label, timeline=True)
    return slide


def _rag_legend(slide, y: int, ref_label: str = '', timeline: bool = False):
    x = int(MARGIN)

    def legend(shape_fn, text: str, width: float):
        nonlocal x
        shape_fn(x, y)
        _add_text(slide, x + Inches(0.26), y - Inches(0.03), Inches(width), Inches(0.25), text, size=9, color=MUTED)
        x += int(Inches(width + 0.4))

    if timeline:
        legend(lambda lx, ly: _add_rect(slide, lx, ly, Inches(0.2), Inches(0.13), PHASE_COLORS[0]), 'Fase', 0.5)
        legend(lambda lx, ly: _add_rect(slide, lx, ly, Inches(0.2), Inches(0.13), SPAN_COLOR), 'Looptijd', 0.7)

        def diamond(filled: bool):
            def draw(lx, ly):
                s = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, lx + Inches(0.02), ly - Inches(0.01), Inches(0.15), Inches(0.15))
                s.shadow.inherit = False
                s.fill.solid()
                s.fill.fore_color.rgb = MILESTONE_COLOR if filled else WHITE
                s.line.color.rgb = MILESTONE_COLOR
            return draw
        legend(diamond(False), 'Mijlpaal gepland', 1.2)
        legend(diamond(True), 'Mijlpaal gehaald', 1.2)
        legend(lambda lx, ly: _add_rect(slide, lx + Inches(0.09), ly - Inches(0.03), Pt(1.5), Inches(0.2), TIMELINE_DEADLINE_COLOR),
               ref_label.capitalize(), 0.8)
    for rag_name, text in (('groen', 'Op koers'), ('oranje', 'Risico'), ('rood', 'Achter'), (None, 'Geen status')):
        color = RAG_COLORS[rag_name] if rag_name else _rag_color(None)
        legend(lambda lx, ly, c=color: _add_rect(slide, lx + Inches(0.06), ly - Inches(0.02), Inches(0.08), Inches(0.18), c),
               text, 0.85)


# ---- slide 1b: alle projecten in één oogopslag --------------------------------

def project_summary(p: dict[str, Any], ref: date) -> dict[str, Any]:
    """Stand van één project voor het overzicht: aantal open en te late
    deliverables, de eerstvolgende verwachte datum (op/na de peildatum) en
    of er überhaupt een planning (een datum) is."""
    open_items = [prod for prod in p['products'] if not prod.get('werkelijkeDatum')]
    open_dates = [d for d in (_parse_iso_date(prod.get('verwachteDatum')) for prod in open_items) if d]
    upcoming = [d for d in open_dates if d >= ref]
    return {
        'products': len(p['products']),
        'open': len(open_items),
        'late': sum(1 for d in open_dates if d < ref),
        'next': min(upcoming) if upcoming else None,
        'planned': bool(project_dates(p)),
    }


def summary_text(summary: dict[str, Any]) -> str:
    if not summary['products'] and not summary['planned']:
        return 'Nog geen planning'
    if summary['products'] and not summary['open']:
        return 'Alles opgeleverd'
    bits = []
    if summary['products']:
        bits.append(f"{summary['open']} open")
    if summary['late']:
        bits.append(f"{summary['late']} te laat")
    if summary['next']:
        bits.append(f"volgende {_fmt_date(summary['next'].isoformat())}")
    elif not summary['planned']:
        bits.append('nog geen datums')
    return ' · '.join(bits)


def overview_grid(n: int) -> tuple[int, int]:
    """(kolommen, rijen) voor het raster van 'Alle projecten'."""
    cols = 1 if n <= 5 else 2 if n <= 12 else 3 if n <= 27 else 4
    return cols, max(1, -(-n // cols))


def _slide_all_projects(prs, projects, ref: date, subtitle: str, links: _Links):
    slide = _blank_slide(prs)
    _slide_header(slide, 'Programma', 'Alle projecten')
    _add_text(slide, MARGIN, Inches(1.48), CONTENT_W, Inches(0.3), subtitle, size=12, color=MUTED)

    n = len(projects)
    cols, rows = overview_grid(n)
    top = int(Inches(1.9))
    bottom = int(SLIDE_H - Inches(0.9))
    gap = int(Inches(0.1 if rows <= 8 else 0.06))
    tile_w = (int(CONTENT_W) - gap * (cols - 1)) // cols
    tile_h = min(int(Inches(0.9)), (bottom - top - gap * (rows - 1)) // rows)
    tile_pt = tile_h / 12700
    name_size = max(7.0, min(13.0, round(tile_pt * 0.3, 1)))
    info_size = max(6.5, name_size - 2)
    two_lines = tile_pt >= (name_size + info_size) * 1.3 + 4
    strip_w = int(Inches(0.07))
    text_w = tile_w - strip_w - int(Inches(0.16))

    for i, p in enumerate(projects):
        proj = p['project']
        code = str(proj.get('code') or '')
        name = str(proj.get('name') or code)
        # Kolomsgewijs vullen, zodat het alfabet van boven naar beneden loopt.
        c, r = divmod(i, rows)
        left = int(MARGIN) + c * (tile_w + gap)
        y = top + r * (tile_h + gap)
        summary = project_summary(p, ref)
        info = (f'{code} · ' if code else '') + summary_text(summary)
        empty = not summary['products'] and not summary['planned']

        tile = _add_rect(slide, left, y, tile_w, tile_h, LIGHT_BG if empty else WHITE, line=True)
        rag = ((proj.get('status') or {}) if isinstance(proj.get('status'), dict) else {}).get('rag')
        strip = _add_rect(slide, left, y, strip_w, tile_h, _rag_color(rag))
        tf = tile.text_frame
        # Altijd afbreken aan: de tekst is al op maat ingekort, en zonder
        # afbreken centreert PowerPoint/LibreOffice de regel.
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = strip_w + int(Inches(0.1))
        tf.margin_right = int(Inches(0.06))
        tf.margin_top = tf.margin_bottom = 0
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        name_chars = _fit_chars(text_w, name_size)
        info_chars = _fit_chars(text_w, info_size)
        run = para.add_run()
        run.font.name = 'Calibri'
        run.font.bold = True
        run.font.size = Pt(name_size)
        run.font.color.rgb = MUTED if empty else DARK
        if two_lines:
            run.text = _truncate(name, name_chars)
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            parts = [(info, None)]
            if summary['late']:
                # 'n te laat' in rood
                before, _, after = info.partition(f"{summary['late']} te laat")
                parts = [(before, None), (f"{summary['late']} te laat", LATE_COLOR), (after, None)]
            used = 0
            for text, color in parts:
                if not text or used >= info_chars:
                    continue
                text = _truncate(text, info_chars - used) if used + len(text) > info_chars else text
                used += len(text)
                r2 = p2.add_run()
                r2.text = text
                r2.font.name = 'Calibri'
                r2.font.size = Pt(info_size)
                r2.font.bold = color is not None
                r2.font.color.rgb = color or MUTED
        else:
            name_part = _truncate(name, max(6, name_chars // 2))
            run.text = name_part
            r2 = para.add_run()
            r2.text = '  ' + _truncate(info, max(4, info_chars - len(name_part) - 2))
            r2.font.name = 'Calibri'
            r2.font.size = Pt(info_size)
            r2.font.color.rgb = LATE_COLOR if summary['late'] else MUTED
        links.add(tile, _project_key(code), f'Naar de voortgang van {name}')
        links.add(strip, _project_key(code), f'Naar de voortgang van {name}')

    _rag_legend(slide, int(SLIDE_H - Inches(0.75)))
    return slide


# ---- slide 2: deliverables ------------------------------------------------

DELIV_COLS = [  # (kop, breedte in inch; None = de rest)
    ('Datum', 1.3), ('Project', 2.9), ('Deliverable', None), ('Type', 1.1), ('Deadline', 1.3), ('Gereed', 0.8),
]


def _deliv_col_widths() -> list[int]:
    fixed = sum(int(Inches(w)) for _, w in DELIV_COLS if w is not None)
    return [int(Inches(w)) if w is not None else int(CONTENT_W) - fixed for _, w in DELIV_COLS]


def _slide_deliverables(prs, entries: list[tuple], titles: dict[str, str], counts: dict[str, int], meta, page: int,
                        subtitle: str, links: _Links):
    slide = _blank_slide(prs)
    _slide_header(slide, 'Programma', 'Deliverables')
    _add_text(slide, MARGIN, Inches(1.48), CONTENT_W, Inches(0.3), subtitle, size=12, color=MUTED)

    widths = _deliv_col_widths()
    row_h = int(Inches(0.32))
    top = int(Inches(1.88))
    _add_rect(slide, MARGIN, top, CONTENT_W, row_h, ACCENT)
    x = int(MARGIN)
    for (head, _), w in zip(DELIV_COLS, widths):
        _add_text(slide, x + Inches(0.08), top, w - Inches(0.1), row_h, head, size=10, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        x += w
    y = top + row_h

    for i, entry in enumerate(entries):
        if entry[0] == 'section':
            key, continued = entry[1], entry[2]
            late = key == 'late'
            _add_rect(slide, MARGIN, y, CONTENT_W, row_h, LATE_BG if late else SECTION_BG)
            text = f'{titles[key]} ({counts[key]})' + (' — vervolg' if continued else '')
            _add_text(slide, MARGIN + Inches(0.08), y, CONTENT_W - Inches(0.16), row_h, text, size=10, bold=True,
                      color=LATE_COLOR if late else ACCENT, anchor=MSO_ANCHOR.MIDDLE)
            y += row_h
            continue
        key, row = entry[1], entry[2]
        late = key == 'late'
        color = LATE_COLOR if late else DARK
        if i % 2 == 0:
            _add_rect(slide, MARGIN, y, CONTENT_W, row_h, LIGHT_BG)
        values = [
            _fmt_date(row['date'].isoformat()) if row['date'] else '—',
            _truncate(row['projectName'] or row['projectCode'], 34),
            _truncate(row['name'], 60),
            row['type'],
            _fmt_date(row['deadline'].isoformat()) if row['deadline'] else '—',
            f"{row['pct']}%",
        ]
        x = int(MARGIN)
        for c, (value, w) in enumerate(zip(values, widths)):
            box = _add_text(slide, x + Inches(0.08), y, w - Inches(0.1), row_h, value, size=10,
                            bold=(late and c == 0), color=color, anchor=MSO_ANCHOR.MIDDLE)
            if c in (1, 2):
                links.add(box, _project_key(row['projectCode']), f'Naar de voortgang van {row["projectName"]}')
            x += w
        y += row_h

    _footer(slide, meta, page, links)
    return slide


def _slide_message(prs, kicker: str, title: str, message: str, meta, page: int, links: _Links):
    slide = _blank_slide(prs)
    _slide_header(slide, kicker, title)
    _add_text(slide, MARGIN, Inches(1.8), CONTENT_W, Inches(0.5), message, size=14, color=MUTED)
    _footer(slide, meta, page, links)
    return slide


# ---- slide 3: bijdrage aan doelen ----------------------------------------

def _names(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(v) for v in value if isinstance(v, str) and v.strip()]


def _slide_goals(prs, projects_page, goal_columns: dict[str, Any], meta, page: int, subtitle: str, links: _Links):
    slide = _blank_slide(prs)
    _slide_header(slide, 'Programma', 'Bijdrage aan doelen')
    _add_text(slide, MARGIN, Inches(1.48), CONTENT_W, Inches(0.3), subtitle, size=12, color=MUTED)

    direct_title = str(goal_columns.get('direct') or 'Draagt direct bij aan')
    top_title = goal_columns.get('top')
    cols = [('Project', Inches(3.0)), (direct_title, None)]
    if top_title:
        cols.append((str(top_title), None))
    flex = (int(CONTENT_W) - int(Inches(3.0))) // (len(cols) - 1)
    widths = [int(Inches(3.0))] + [flex] * (len(cols) - 1)

    head_h = int(Inches(0.36))
    top = int(Inches(1.88))
    _add_rect(slide, MARGIN, top, CONTENT_W, head_h, ACCENT)
    x = int(MARGIN)
    for (head, _), w in zip(cols, widths):
        _add_text(slide, x + Inches(0.08), top, w - Inches(0.1), head_h, _truncate(head, 60), size=10, bold=True,
                  color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        x += w
    y = top + head_h
    row_h = (int(SLIDE_H - Inches(0.65)) - y) // GOAL_ROWS_PER_SLIDE

    for i, p in enumerate(projects_page):
        proj = p['project']
        code = str(proj.get('code') or '')
        name = str(proj.get('name') or code)
        if i % 2 == 1:
            _add_rect(slide, MARGIN, y, CONTENT_W, row_h, LIGHT_BG)
        label = _add_text(slide, MARGIN + Inches(0.08), y + Inches(0.06), widths[0] - Inches(0.16), row_h - Inches(0.1),
                          _truncate(name, 60), size=11, bold=True, color=DARK)
        links.add(label, _project_key(code), f'Naar de voortgang van {name}')
        goals = p['goals']
        cells = [_names(goals.get('direct'))]
        if top_title:
            cells.append(_names(goals.get('top')))
        x = int(MARGIN) + widths[0]
        for names, w in zip(cells, widths[1:]):
            text = ' · '.join(names) if names else '—'
            _add_text(slide, x + Inches(0.08), y + Inches(0.06), w - Inches(0.16), row_h - Inches(0.1),
                      _truncate(text, 200), size=10, color=DARK if names else MUTED)
            x += w
        y += row_h

    _footer(slide, meta, page, links)
    return slide


# ---- alles samen ------------------------------------------------------------

def _chunks(items: list[Any], size: int) -> list[list[Any]]:
    return [items[i:i + size] for i in range(0, len(items), size)] or [[]]


def planned_slide_count(projects: list[dict[str, Any]], ref: date, window_months: int) -> int:
    deliv_pages = paginate_deliverables(classify_deliverables(projects, ref, window_months))
    return (2 + max(1, len(deliv_pages))
            + len(_chunks(projects, GOAL_ROWS_PER_SLIDE)) + len(projects))


def build_program_pptx(data: Any, options: Any, meta: Any) -> bytes:
    meta = _clean_data(meta) if isinstance(meta, dict) else {}
    projects = sort_projects(_projects_from(_clean_data(data)))
    ref, window_months = parse_options(options if isinstance(options, dict) else None, meta)
    goal_columns = (data or {}).get('goalColumns') if isinstance((data or {}).get('goalColumns'), dict) else {}
    goal_columns = _clean_data(goal_columns)

    sections = classify_deliverables(projects, ref, window_months)
    deliv_pages = paginate_deliverables(sections)
    total = planned_slide_count(projects, ref, window_months)
    if total > MAX_SLIDES:
        raise ProgramPptxError(
            f'Deze keuze levert {total} slides op; het maximum is {MAX_SLIDES}. Kies minder projecten.'
        )

    exported = _parse_iso_date(meta.get('exportedAt'))
    ref_label = 'vandaag' if exported is None or exported == ref else 'peildatum'
    prs = _new_presentation()
    links = _Links()
    page = 1

    # 1. Tijdlijnen (alleen projecten met datums, altijd één slide) -- Home
    with_dates = [p for p in projects if project_dates(p)]
    all_dates = [d for p in with_dates for d in project_dates(p)]
    subtitle = f'{len(with_dates)} van {len(projects)} project(en) met een planning · {ref_label} {_fmt_date(ref.isoformat())}'
    if len(with_dates) < len(projects):
        subtitle += ' · alle projecten op de volgende slide'
    slide = _slide_timeline(prs, with_dates, all_dates, ref, subtitle, links, ref_label)
    _footer(slide, meta, page, None)
    links.slides[HOME_KEY] = slide
    page += 1

    # 1b. Alle projecten in één oogopslag (ook zonder planning)
    n_empty = sum(1 for p in projects if not p['products'] and not project_dates(p))
    subtitle = f'{len(projects)} project(en)' + (f' · {n_empty} nog zonder planning' if n_empty else '') + \
        ' · klik op een project voor de voortgang'
    slide = _slide_all_projects(prs, projects, ref, subtitle, links)
    _footer(slide, meta, page, links)
    page += 1

    # 2. Deliverables
    titles = section_titles(ref, window_months)
    counts = {k: len(v) for k, v in sections.items()}
    open_total = sum(counts.values())
    if deliv_pages:
        for i, entries in enumerate(deliv_pages):
            bits = [f'{open_total} openstaande deliverable(s)']
            if counts['late']:
                bits.append(f'{counts["late"]} te laat')
            subtitle = ' · '.join(bits) + _page_suffix(i, len(deliv_pages))
            _slide_deliverables(prs, entries, titles, counts, meta, page, subtitle, links)
            page += 1
    else:
        _slide_message(prs, 'Programma', 'Deliverables',
                       'Er zijn geen openstaande deliverables in de gekozen projecten.', meta, page, links)
        page += 1

    # 3. Bijdrage aan doelen
    goal_pages = _chunks(projects, GOAL_ROWS_PER_SLIDE)
    for i, chunk in enumerate(goal_pages):
        subtitle = 'Waar dragen de projecten aan bij in de doelenboom?' + _page_suffix(i, len(goal_pages))
        _slide_goals(prs, chunk, goal_columns, meta, page, subtitle, links)
        page += 1

    # 4. Voortgang per project (alfabetisch). 'Vandaag' op de projecttijdlijn
    # is de peildatum (meta.referenceDate, zie _today_iso in project_pptx.py),
    # zodat de hele presentatie dezelfde datum gebruikt.
    project_meta = {**meta, 'referenceDate': ref.isoformat()}
    for p in projects:
        proj = p['project']
        slide = _slide_overview(prs, proj, p['products'], project_meta, page)
        _home_button(slide, links)
        links.slides[_project_key(str(proj.get('code') or ''))] = slide
        page += 1

    links.resolve()
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
